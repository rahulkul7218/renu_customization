import frappe
from frappe import _
from frappe.utils import flt
import json
import base64
from io import BytesIO
import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


def execute(filters=None):
    if not filters:
        filters = {}

    columns = get_columns()
    data = get_data(filters)

    return columns, data


def get_columns():
    return [
        {"label": "name", "fieldname": "name", "fieldtype": "Data", "hidden": 1},
        _("Supplier PO No.") + ":Link/Purchase Order:150",
        _("Supplier PO Date") + ":Date:120",
        {
            "label": _("Sr.No."),
            "fieldname": "sr_no",
            "fieldtype": "Int",
            "width": 70,
            "disable_total": 1
        },
        _("Linked SO No.") + ":Link/Sales Order:150",
        _("SO Date") + ":Date:120",
        _("Supplier Code") + ":Link/Supplier:150",
        _("Supplier Name") + ":Data:180",
        _("Item Code") + ":Link/Item:120",
        _("Item Name") + ":Data:180",
        _("Order Quantity") + ":Float:130",
        _("Delivered Qty") + ":Float:120",
        _("Open Qty") + ":Float:120",
        _("Item Rate") + ":Float:120",
        _("Currency") + ":Link/Currency:100",
        {
            "label": "Exchange Rate",
            "fieldname": "exchange_rate",
            "fieldtype": "Float",
            "disable_total": 1
        },
        _("Total Net Amount (INR)") + ":Float:180",
        _("Delivered Net Total") + ":Float:180",
        _("Balance Net Total") + ":Float:170",
        _("Delivery Date") + ":Date:120",
    ]


def get_conditions(filters):
    if not filters:
        filters = {}
    conditions = ""

    # -------- Status filter (multiple selection) --------
    status = filters.get("status")
    if status:
        if isinstance(status, str):
            status_list = [s.strip() for s in status.split(",") if s.strip()]
        else:
            status_list = status

        if status_list:
            filters["status"] = tuple(status_list)
            conditions += " AND po.status IN %(status)s"

    # -------- Other filters --------
    if filters.get("po_no"):
        conditions += " AND po.name = %(po_no)s"

    if filters.get("supplier_name"):
        filters["supplier_name"] = f"%{filters['supplier_name']}%"
        conditions += " AND po.supplier_name LIKE %(supplier_name)s"

    if filters.get("item_code"):
        conditions += " AND poi.item_code = %(item_code)s"

    if filters.get("currency"):
        conditions += " AND po.currency = %(currency)s"

    if filters.get("from_date"):
        conditions += " AND po.transaction_date >= %(from_date)s"

    if filters.get("to_date"):
        conditions += " AND po.transaction_date <= %(to_date)s"

    return conditions


def get_data(filters):
    if not filters:
        filters = {}
    conditions = get_conditions(filters)

    sql = f"""
        SELECT
            poi.name AS name,
            po.name AS supplier_po_no,
            po.transaction_date AS supplier_po_date,
            ROW_NUMBER() OVER (PARTITION BY po.name ORDER BY poi.idx) AS sr_no,
            poi.sales_order AS linked_so_no,
            (SELECT transaction_date FROM `tabSales Order` WHERE name = poi.sales_order LIMIT 1) AS so_date,
            po.supplier AS supplier_code,
            po.supplier_name AS supplier_name,
            poi.item_code AS item_code,
            poi.item_name AS item_name,
            poi.qty AS order_quantity,
            poi.received_qty AS delivered_qty,
            (poi.qty - IFNULL(poi.received_qty, 0)) AS open_qty,
            poi.rate AS item_rate,
            po.currency AS currency,
            po.conversion_rate AS exchange_rate,
            (poi.qty * poi.base_rate) AS `total_net_amount_(inr)`,
            (poi.received_qty * poi.base_rate) AS delivered_net_total,
            ((poi.qty - IFNULL(poi.received_qty, 0)) * poi.base_rate) AS balance_net_total,
            poi.schedule_date AS delivery_date

        FROM `tabPurchase Order` po
        INNER JOIN `tabPurchase Order Item` poi ON poi.parent = po.name
       
        WHERE 1 = 1
        AND po.docstatus = 1
        
        {conditions}

        ORDER BY
         po.transaction_date ASC,
         po.name ASC,
         poi.idx ASC
    """

    return frappe.db.sql(sql, filters, as_list=True)


@frappe.whitelist()
def download_xlsx(filters=None, include_filters=1):
    if isinstance(filters, str):
        try:
            filters = frappe.parse_json(filters)
        except Exception:
            filters = json.loads(filters)

    if not isinstance(filters, dict):
        filters = {}

    include_filters = frappe.utils.cint(filters.get("include_filters", include_filters))

    if include_filters:
        columns, data = execute(filters)
    else:
        columns, data = execute({})
        
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Purchase Order Report"

    ws["A1"].value = "Report Name"
    ws["A1"].font = Font(bold=True)
    ws["B1"].value = "Purchase Order Report"

    ws["A2"].value = "Generated On"
    ws["A2"].font = Font(bold=True)
    ws["B2"].value = frappe.utils.now_datetime().strftime("%Y-%m-%d %H:%M:%S")

    ws["A3"].value = "Generated By"
    ws["A3"].font = Font(bold=True)
    full_name = frappe.db.get_value("User", frappe.session.user, "full_name")
    ws["B3"].value = full_name or frappe.session.user

    row_idx = 4
    if filters:
        for key, val in filters.items():
            if not val or key in ["include_filters", "report_name", "current_datetime"]:
                continue
            
            label = frappe.unscrub(key)
            if isinstance(val, (list, tuple)):
                val = ", ".join([str(v) for v in val])
            
            ws.cell(row=row_idx, column=1, value=label).font = Font(bold=True)
            ws.cell(row=row_idx, column=2, value=str(val))
            row_idx += 1
    
    row_idx += 1

    keep_indices = []
    filtered_columns = []
    for i, col in enumerate(columns):
        if isinstance(col, dict) and col.get("hidden"):
            continue
        keep_indices.append(i)
        filtered_columns.append(col)
    
    columns = filtered_columns
    data = [[row[i] for i in keep_indices] for row in data]

    header_row = row_idx
    for idx, col in enumerate(columns, start=1):
        label = col["label"] if isinstance(col, dict) else col.split(":")[0]
        c = ws.cell(row=header_row, column=idx, value=label)
        c.font = Font(bold=True)
        c.alignment = Alignment(horizontal="center")

    row_idx += 1

    numeric_index_map = {}
    no_total_index_set = set()
    
    for i, col in enumerate(columns):
        label = col["label"] if isinstance(col, dict) else col.split(":")[0]
        f_type = col.get("fieldtype") if isinstance(col, dict) else (col.split(":")[1] if ":" in col else "")
        
        if f_type in ["Float", "Int", "Currency", "Percent"]:
            numeric_index_map[i] = True
            
        if label in ["Supplier PO Date", "SO Date", "Delivery Date", "Item Rate", "Currency", "Exchange Rate"]:
            no_total_index_set.add(i)

    for row in data:
        for col_idx, value in enumerate(row, start=1):
            cell = ws.cell(row=row_idx, column=col_idx)
            
            col_def = columns[col_idx - 1]
            label = col_def["label"] if isinstance(col_def, dict) else col_def.split(":")[0]
            fieldname = label.lower().replace(" ", "_")

            if fieldname in ("sr.no.", "sr_no", "sr_no.", "sr.no"):
                cell.value = int(value) if value else 0
                cell.number_format = "0"
                cell.alignment = Alignment(horizontal="center")
                continue

            if (col_idx - 1) in numeric_index_map and value not in (None, ""):
                try:
                    cell.value = float(value)
                    cell.number_format = "#,##0.00"
                except:
                    cell.value = value
                cell.alignment = Alignment(horizontal="right")
            else:
                cell.value = value
                cell.alignment = Alignment(horizontal="left")

        row_idx += 1

    total_row = row_idx
    for col_idx in range(1, len(columns) + 1):
        cell = ws.cell(row=total_row, column=col_idx)
        cell.fill = PatternFill(start_color="D3D3D3", fill_type="solid")
        cell.font = Font(bold=True)

        if col_idx == 1:
            cell.value = "Total"
            cell.alignment = Alignment(horizontal="left")
            continue

        field_idx = col_idx - 1
        if field_idx in numeric_index_map and field_idx not in no_total_index_set:
            try:
                total_val = sum(flt(r[field_idx]) for r in data)
            except:
                total_val = 0
            cell.value = total_val
            cell.number_format = "#,##0.00"
            cell.alignment = Alignment(horizontal="right")
        else:
            cell.value = ""

    for i in range(1, len(columns) + 1):
        ws.column_dimensions[get_column_letter(i)].width = 22

    out = BytesIO()
    wb.save(out)
    out.seek(0)

    return base64.b64encode(out.read()).decode()
