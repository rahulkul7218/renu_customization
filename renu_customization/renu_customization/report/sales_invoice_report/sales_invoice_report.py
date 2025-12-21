from warnings import filters
import frappe
 
def execute(filters=None):
    if not filters:
        filters = {}
 
    conditions = " WHERE 1=1 "
 
    status = filters.get("status")
 
    if status:
        # From report UI it comes as a list, e.g. ["Draft", "To Deliver"]
        if isinstance(status, str):
            status_list = [s.strip() for s in status.split(",") if s.strip()]
        else:
            status_list = status
 
        if status_list:
            filters["status"] = tuple(status_list)
            conditions += " AND si.status IN %(status)s"
       
    # Invoice ID Filter
    if filters.get("invoice_id"):
        conditions += " AND si.name = %(invoice_id)s"
 
    # Invoice Date From
    if filters.get("from_date"):
        conditions += " AND si.posting_date >= %(from_date)s"
 
    # Invoice Date To
    if filters.get("to_date"):
        conditions += " AND si.posting_date <= %(to_date)s"
 
    # Customer Name
    if filters.get("customer_name"):
        conditions += " AND si.customer_name LIKE %(customer_name)s"
 
    # Item Code
    if filters.get("item_code"):
        conditions += " AND sii.item_code = %(item_code)s"
 
    # Currency
    if filters.get("currency"):
        conditions += " AND si.currency = %(currency)s"
 
   
 
    # City
    if filters.get("city"):
        conditions += " AND ad.city LIKE %(city)s"
 
    # State
    if filters.get("state"):
        conditions += " AND ad.state LIKE %(state)s"
 
    # Country
    if filters.get("country"):
        conditions += " AND ad.country LIKE %(country)s"
 
 
    # ---------------------------
    # Define columns (REQUIRED)
    # ---------------------------
    columns = [
        {"label": "Invoice ID", "fieldname": "invoice_id", "fieldtype": "Link", "options": "Sales Invoice", "width": 120},
        {"label": "Invoice Date", "fieldname": "invoice_date", "fieldtype": "Date", "width": 120},
        {"label": "Sr.No.", "fieldname": "sr_no", "fieldtype": "Int", "width": 70, "disable_total": 1},
        {"label": "Customer PO No.", "fieldname": "po_no", "fieldtype": "Data", "width": 150},
        {"label": "Customer PO Date", "fieldname": "po_date", "fieldtype": "Date", "width": 150},
 
        {"label": "Order ID No.", "fieldname": "order_id_no", "fieldtype": "Link", "options": "Sales Order", "width": 150},
        {"label": "Order ID Date", "fieldname": "order_id_date", "fieldtype": "Date", "width": 150},
 
        {"label": "Customer Code", "fieldname": "customer_code", "fieldtype": "Data", "width": 150},
        {"label": "Customer Name", "fieldname": "customer_name", "fieldtype": "Data", "width": 150},
 
        # {"label": "Party Item Code", "fieldname": "party_item_code", "fieldtype": "Data", "width": 150},
        {"label": "Item Code", "fieldname": "item_code", "fieldtype": "Link", "options": "Item", "width": 120},
        {"label": "Item Name", "fieldname": "item_name", "fieldtype": "Data", "width": 140},
        {"label": "Item Description", "fieldname": "description", "fieldtype": "Data", "width": 220},
        {"label": "Qty", "fieldname": "qty", "fieldtype": "Float", "width": 80},
 
        {"label": "Item Rate", "fieldname": "item_rate", "fieldtype": "Float", "width": 120},
        {"label": "Amount", "fieldname": "amount", "fieldtype": "Float", "width": 120},
        {"label": "Currency", "fieldname": "currency", "fieldtype": "Data", "width": 120},
        {"label": "Exchange Rate", "fieldname": "exchange_rate", "fieldtype": "Float", "width": 140, "disable_total": 1},
        {"label": "Amount (INR)", "fieldname": "base_amount", "fieldtype": "Float", "width": 120},
 
        {"label": "Delivery Date", "fieldname": "delivery_date", "fieldtype": "Date", "width": 120},
 
        {"label": "Business Region Name", "fieldname": "business_region_name", "fieldtype": "Data", "width": 190},
        {"label": "City", "fieldname": "city", "fieldtype": "Data", "width": 120},
        {"label": "State", "fieldname": "state", "fieldtype": "Data", "width": 120},
        {"label": "Country", "fieldname": "country", "fieldtype": "Data", "width": 120},
        {"label": "Sales Person", "fieldname": "sales_person", "fieldtype": "Data", "width": 150},
 
        {"label": "Domestic/Export", "fieldname": "dom_exp", "fieldtype": "Data", "width": 150},
        {"label": "Item Purchase Rate", "fieldname": "item_purchase_rate", "fieldtype": "Float", "width": 180},
        {"label": "OldNewFlg", "fieldname": "old_new_flg", "fieldtype": "Data", "width": 100},
        {"label": "Business Activity", "fieldname": "business_activity", "fieldtype": "Data", "width": 150},
        {"label": "Business Vertical", "fieldname": "business_vertical", "fieldtype": "Data", "width": 150},
    ]
 
 
    # ---------------------
    # SQL Query
    # ---------------------
    query = f"""
        SELECT
            si.name AS invoice_id,
            si.posting_date AS invoice_date,
           ROW_NUMBER() OVER (PARTITION BY si.name ORDER BY sii.idx) AS sr_no,
 
            si.po_no AS po_no,
 
            si.po_date AS po_date,
 
            so.name AS order_id_no,
 
            so.transaction_date AS order_id_date,
 
            c.customer_code AS customer_code,
 
            si.customer_name AS customer_name,
 
            sii.item_code AS item_code,
 
            sii.item_name AS item_name,
 
            REGEXP_REPLACE(sii.description, '<[^>]*>', '') AS description,
           
            sii.qty AS qty,
 
            sii.rate AS item_rate,
            sii.amount AS amount,
            si.currency AS currency,
            si.conversion_rate AS exchange_rate,
            sii.base_amount AS base_amount,
 
            dn.posting_date AS delivery_date,
 
            c.business_region_name AS business_region_name,
            ad.city AS city,
            ad.state AS state,
            ad.country AS country,
            st.sales_person AS sales_person,
 
            CASE WHEN ad.country = 'India' THEN 'Domestic' ELSE 'Export' END AS dom_exp,
 
            (
            SELECT
                IFNULL(sle.incoming_rate, 0)
            FROM `tabStock Ledger Entry` sle
            WHERE sle.item_code = sii.item_code
            AND sle.actual_qty > 0           -- Only incoming entries
            ORDER BY sle.posting_date DESC, sle.posting_time DESC
            LIMIT 1
        ) AS item_purchase_rate,
 
            '' AS old_new_flg,
            '' AS business_activity,
            '' AS business_vertical
 
        FROM `tabSales Invoice` si
 
        JOIN `tabSales Invoice Item` sii ON sii.parent = si.name
        # AND (si.amended_from IS NULL OR si.name = (
        #     SELECT MAX(name)
        #     FROM `tabSales Invoice`
        #     WHERE name LIKE CONCAT(SUBSTRING_INDEX(si.name, '-', 1), '%%')
        # ))
 
         AND si.name = (
            SELECT MAX(name)
            FROM `tabSales Invoice`
            WHERE name LIKE CONCAT(SUBSTRING_INDEX(si.name, '-', 1), '%%')
            AND docstatus != 2
        )
 
        LEFT JOIN `tabSales Order Item` soi ON soi.name = sii.so_detail
        LEFT JOIN `tabSales Order` so ON so.name = soi.parent
       
 
        # LEFT JOIN `tabDynamic Link` dl ON dl.link_name = si.customer
        #     AND dl.link_doctype = 'Customer'
        #     AND dl.parenttype = 'Address'
 
        # LEFT JOIN `tabAddress` ad ON ad.name = dl.parent

        LEFT JOIN `tabAddress` ad ON ad.name = si.customer_address
 
 
        LEFT JOIN `tabCustomer` c ON c.name = si.customer
 
        LEFT JOIN `tabSales Team` st ON st.parent = si.name

        
        LEFT JOIN `tabItem` it ON it.name = sii.item_code
 
        LEFT JOIN `tabItem Price` ip ON ip.item_code = sii.item_code
            AND ip.buying = 1 AND ip.selling = 0
 
        LEFT JOIN `tabDelivery Note Item` dni ON sii.dn_detail = dni.name
        LEFT JOIN `tabDelivery Note` dn ON dn.name = dni.parent
 
        {conditions}
        AND it.is_stock_item = 1
        
        ORDER BY si.posting_date ASC, si.name ASC,sii.idx ASC
    """
 
    data = frappe.db.sql(query, filters, as_dict=1)
    return columns, data
 
 
import frappe
from frappe.utils import flt
from frappe import _
import base64
from io import BytesIO
import openpyxl
from openpyxl.utils import get_column_letter
from openpyxl.styles import Alignment, Font, PatternFill
import json
 
@frappe.whitelist()
def download_xlsx(filters=None):
 
    # ---------- PARSE FILTERS ----------
    if isinstance(filters, str):
        try:
            filters = frappe.parse_json(filters)
        except Exception:
            filters = json.loads(filters)
 
    if not isinstance(filters, dict):
        filters = {}
 
    # read include_filters flag
    include_filters = frappe.utils.cint(filters.get("include_filters", 1))
 
    # remove include_filters from actual filter list
    actual_filters = {k: v for k, v in filters.items() if k != "include_filters"}
 
    # ---------- GET REPORT DATA ----------
    columns, data = execute(filters)
 
 
 
 
   
 
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sales Invoice Report"
 
    # ---------- HEADER ----------
    ws["A1"].value = "Report Name"
    ws["A1"].font = Font(bold=True)
    ws["B1"].value = "Sales Invoice Report"
 
    ws["A2"].value = "Generated On"
    ws["A2"].font = Font(bold=True)
    ws["B2"].value = frappe.utils.now_datetime().strftime("%Y-%m-%d %H:%M:%S")
 
    ws["A3"].value = "Generated By"
    ws["A3"].font = Font(bold=True)
    full_name = frappe.db.get_value("User", frappe.session.user, "full_name")
    ws["B3"].value = full_name or frappe.session.user
   
    ws.append([])
    row_idx = 5
 
    # ---------- FILTERS SECTION ----------
 
 
    # ---------- COLUMN HEADERS ----------
    for idx, col in enumerate(columns, start=1):
        c = ws.cell(row=row_idx, column=idx, value=col["label"])
        c.font = Font(bold=True)
        c.alignment = Alignment(horizontal="center")
 
    row_idx += 1
 
    # numeric fields
    numeric_fields = {   "qty", "item_rate",
         "exchange_rate", "item_purchase_rate" ,"amount", "base_amount" }
 
    # Skip total for selected numeric columns
    no_total_fields = {
        #  "qty",
        #  "item_rate",
         "exchange_rate",
        #  "item_purchase_rate"
    }
 
    # ---------- DATA ROWS ----------
    for row in data:
        for col_idx, col in enumerate(columns, start=1):
            fieldname = col["fieldname"]
            value = row.get(fieldname)
            cell = ws.cell(row=row_idx, column=col_idx)
 
            if fieldname == "sr_no":
                cell.value = int(value) if value else 0
                cell.number_format = "0"
                cell.alignment = Alignment(horizontal="center")
 
            elif fieldname in numeric_fields:
                cell.value = flt(value or 0)
                cell.number_format = "#,##0.00"
                cell.alignment = Alignment(horizontal="right")
 
            else:
                if isinstance(value, (list, tuple, set, dict)):
                    value = json.dumps(value)
                cell.value = value
                cell.alignment = Alignment(horizontal="left")
 
        row_idx += 1
 
    # ---------- TOTAL ROW ----------
    total_row = row_idx
    for col_idx, col in enumerate(columns, start=1):
        fieldname = col["fieldname"]
        cell = ws.cell(total_row, col_idx)
 
        cell.fill = PatternFill(start_color="D3D3D3", fill_type="solid")
        cell.font = Font(bold=True)
 
        if col_idx == 1:
            cell.value = "Total"
            cell.alignment = Alignment(horizontal="left")
            continue
       
        # Skip total for selected numeric columns
        if fieldname in no_total_fields:
             cell.value = ""
             continue
 
        if fieldname in numeric_fields:
            total_val = sum(flt(d.get(fieldname) or 0) for d in data)
            cell.value = total_val
            cell.number_format = "#,##0.00"
            cell.alignment = Alignment(horizontal="right")
        else:
            cell.value = ""
 
    # ---------- COLUMN WIDTH ----------
    for idx in range(1, len(columns) + 1):
        ws.column_dimensions[get_column_letter(idx)].width = 22
 
    # ---------- OUTPUT ----------
    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return base64.b64encode(output.read()).decode()
 
 