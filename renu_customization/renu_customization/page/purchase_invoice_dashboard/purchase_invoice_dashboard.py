import frappe
from frappe import _
from frappe.utils import flt
import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
from io import BytesIO
import base64
import json

@frappe.whitelist()
def export_to_pdf(html):
	frappe.response.filename = f"Purchase_Invoice_Dashboard_{frappe.utils.nowdate()}.pdf"
	frappe.response.type = "binary"
	frappe.response.filecontent = frappe.utils.pdf.get_pdf(html, {"orientation": "Landscape"})

def prepare_filters(filters):
    if not filters:
        filters = {}
    elif isinstance(filters, str):
        filters = frappe.parse_json(filters)
    
    # Handle DateRange from JS
    if filters.get("date_range"):
        date_range = filters.get("date_range")
        if isinstance(date_range, list) and len(date_range) == 2:
            filters["from_date"] = date_range[0]
            filters["to_date"] = date_range[1]
    
    return frappe._dict(filters)

@frappe.whitelist()
def get_dashboard_data(filters=None):
    filters = prepare_filters(filters)
    
    # Handle Fiscal Year
    if filters.get("fiscal_year"):
        fy = frappe.get_doc("Fiscal Year", filters.fiscal_year)
        if fy:
            filters["from_date"] = fy.year_start_date
            filters["to_date"] = fy.year_end_date

    conditions = " WHERE pi.docstatus = 1 " # Only submitted invoices
    
    if filters.get("from_date"):
        conditions += " AND pi.posting_date >= %(from_date)s "
    if filters.get("to_date"):
        conditions += " AND pi.posting_date <= %(to_date)s "
    if filters.get("supplier"):
        conditions += " AND pi.supplier = %(supplier)s "
    if filters.get("supplier_group"):
        conditions += " AND s.supplier_group = %(supplier_group)s "
    if filters.get("item_code"):
        conditions += " AND pii.item_code = %(item_code)s "
    if filters.get("status"):
        # MultiSelect status handling
        status_filter = filters.get("status")
        if isinstance(status_filter, str):
            status_filter = [s.strip() for s in status_filter.split(",")]
        filters["status_list"] = tuple(status_filter)
        conditions += " AND pi.status IN %(status_list)s "

    query = f"""
        SELECT
            pi.name AS invoice_id,
            pi.posting_date AS invoice_date,
            pi.supplier AS supplier,
            pi.supplier_name AS supplier_name,
            pi.status AS status,
            pi.is_return AS is_return,
            pii.item_code AS item_code,
            pii.item_name AS item_name,
            pii.qty AS qty,
            pii.base_net_amount AS base_amount,
            pi.paid_amount AS paid_amount,
            pi.outstanding_amount AS outstanding_amount,
            ad.country AS country,
            s.supplier_group AS supplier_group
        FROM `tabPurchase Invoice` pi
        JOIN `tabPurchase Invoice Item` pii ON pii.parent = pi.name
        LEFT JOIN `tabSupplier` s ON s.name = pi.supplier
        LEFT JOIN `tabAddress` ad ON ad.name = pi.supplier_address
        {conditions}
        ORDER BY pi.posting_date DESC, pi.name DESC
    """
    
    raw_data = frappe.db.sql(query, filters, as_dict=1)
    
    if not raw_data:
        return {
            "summary": [], "charts": {}, "results": [], "columns": []
        }

    # Calculate Summaries (KPIs)
    total_pur = 0
    total_paid = 0
    total_unpaid = 0
    ret_pur = 0

    for row in raw_data:
        amt = flt(row.get("base_amount"))
        paid = flt(row.get("paid_amount"))
        unpaid = flt(row.get("outstanding_amount"))
        
        if row.get("is_return"):
            ret_pur += amt
            total_pur -= amt 
        else:
            total_pur += amt
            total_paid += paid
            total_unpaid += unpaid

    report_summary = [
        {"label": _("Total Purchase"), "value": total_pur, "indicator": "blue", "fieldtype": "Currency", "currency": "INR"},
        {"label": _("Total Paid"), "value": total_paid, "indicator": "green", "fieldtype": "Currency", "currency": "INR"},
        {"label": _("Total Unpaid"), "value": total_unpaid, "indicator": "orange", "fieldtype": "Currency", "currency": "INR"},
        {"label": _("Debit Notes"), "value": ret_pur, "indicator": "red", "fieldtype": "Currency", "currency": "INR"}
    ]

    sup_purchase = {}
    prod_purchase = {}
    
    for row in raw_data:
        amt = flt(row.get("base_amount"))
        if row.get("is_return"): amt = -amt
        
        sup = row.get("supplier_name") or row.get("supplier") or ""
        sup_purchase[sup] = sup_purchase.get(sup, 0) + amt
        
        prod = row.get("item_name") or row.get("item_code") or ""
        prod_purchase[prod] = prod_purchase.get(prod, 0) + amt

    def get_chart_def(title, data_dict, limit=10):
        sorted_items = sorted(data_dict.items(), key=lambda x: x[1], reverse=True)
        top_items = sorted_items[:limit]
        return {
            "title": title,
            "data": {
                "labels": [x[0] for x in top_items],
                "datasets": [{"name": title, "values": [flt(x[1], 2) for x in top_items]}]
            },
            "type": "donut",
            "height": 300,
            "colors": ['#3498db', '#e74c3c', '#2ecc71', '#f1c40f', '#9b59b6', '#34495e', '#ecf0f1', '#1abc9c', '#d35400', '#7f8c8d']
        }

    return {
        "summary": report_summary,
        "charts": {
            "top_10_suppliers": get_chart_def("Top 10 Suppliers by Value", sup_purchase),
            "top_10_products": get_chart_def("Top 10 Products by Value", prod_purchase)
        },
        "results": raw_data,
        "columns": [] # JS will define columns for display
    }

@frappe.whitelist()
def export_to_excel(filters=None):
    filters = prepare_filters(filters)
    dashboard_data = get_dashboard_data(filters)
    data = dashboard_data.get("results")
    summary = dashboard_data.get("summary")
    charts = dashboard_data.get("charts")
    
    if not data:
        return None

    wb = openpyxl.Workbook()
    
    # Sheet 1: Dashboard Overview
    ws_overview = wb.active
    ws_overview.title = "Dashboard Overview"
    
    # Sheet 2: Month-Wise Purchase
    ws_months = wb.create_sheet("Month-Wise Purchase")
    
    # Sheet 3: Purchase Invoices List
    ws_list = wb.create_sheet("Purchase Invoices List")
    
    # Styling Helpers
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="2c3e50", fill_type="solid")
    title_font = Font(bold=True, size=14)
    section_font = Font(bold=True, size=12)
    thin_side = Side(style='thin')
    table_border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)
    zebra_fill = PatternFill(start_color="f8f9fa", fill_type="solid")

    row_idx = 1
    
    # -------------------------------------------------------------------------
    # Sheet 1: Dashboard Overview
    # -------------------------------------------------------------------------
    ws_overview.cell(row=row_idx, column=1, value="Purchase Invoice Dashboard Overview").font = title_font
    ws_overview.cell(row=row_idx, column=5, value="Generated On: " + frappe.utils.now_datetime().strftime("%Y-%m-%d %H:%M:%S"))
    row_idx += 2
 
    # 1. Summary Section
    ws_overview.cell(row=row_idx, column=1, value="1. Purchase Summary (Million INR)").font = section_font
    row_idx += 1
    
    summary_start_row = row_idx
    colors = {"blue": "3498db", "green": "2ecc71", "orange": "e67e22", "purple": "9b59b6", "red": "e74c3c"}
 
    for i, s in enumerate(summary):
        r = summary_start_row + (i // 4) * 3
        c = 1 + (i % 4) * 2
        cell_l = ws_overview.cell(row=r, column=c, value=s.get('label'))
        cell_l.font = Font(bold=True, color="FFFFFF")
        bg_color = colors.get(s.get('indicator', 'blue').lower(), "3498db")
        cell_l.fill = PatternFill(start_color=bg_color, fill_type="solid")
        cell_l.alignment = Alignment(horizontal="center")
        
        val = flt(s.get('value')) / 1000000
        cell_v = ws_overview.cell(row=r+1, column=c, value=val)
        cell_v.font = Font(bold=True, size=12)
        cell_v.number_format = '"₹ "#,##0.00" M"'
        cell_v.alignment = Alignment(horizontal="center")
        cell_v.border = Border(left=Side(style='medium', color=bg_color), right=Side(style='medium', color=bg_color), bottom=Side(style='medium', color=bg_color))
        ws_overview.merge_cells(start_row=r, start_column=c, end_row=r, end_column=c+1)
        ws_overview.merge_cells(start_row=r+1, start_column=c, end_row=r+1, end_column=c+1)
 
    row_idx = summary_start_row + 6
    row_idx += 1
 
    def write_chart_section(title, chart_data):
        nonlocal row_idx
        ws_overview.cell(row=row_idx, column=1, value=title).font = section_font
        row_idx += 1
        headers = ["Category", "Amount (M)", "Share %"]
        for idx, h in enumerate(headers, start=1):
            cell = ws_overview.cell(row=row_idx, column=idx, value=h)
            cell.font = header_font
            cell.fill = PatternFill(start_color="34495e", fill_type="solid")
            cell.border = table_border
            cell.alignment = Alignment(horizontal="center")
        row_idx += 1
        labels = chart_data.get("data", {}).get("labels", [])
        values = chart_data.get("data", {}).get("datasets", [{}])[0].get("values", [])
        total_val = sum(values) if values else 1
        for i in range(len(labels)):
            row_fill = zebra_fill if i % 2 == 0 else None
            c1 = ws_overview.cell(row=row_idx, column=1, value=labels[i]); c1.border = table_border
            if row_fill: c1.fill = row_fill
            val_m = flt(values[i]) / 1000000
            c2 = ws_overview.cell(row=row_idx, column=2, value=val_m); c2.number_format = '"₹ "#,##0.00" M"'; c2.border = table_border; c2.alignment = Alignment(horizontal="right")
            if row_fill: c2.fill = row_fill
            share = (values[i] / total_val) if total_val else 0
            c3 = ws_overview.cell(row=row_idx, column=3, value=share); c3.number_format = "0.00%"; c3.border = table_border; c3.alignment = Alignment(horizontal="center")
            if row_fill: c3.fill = row_fill
            row_idx += 1
        row_idx += 2
        
    if charts.get("top_10_suppliers", {}).get("data", {}).get("labels"):
        write_chart_section("2. Top 10 Suppliers", charts["top_10_suppliers"])
    if charts.get("top_10_products", {}).get("data", {}).get("labels"):
        write_chart_section("3. Top 10 Products", charts["top_10_products"])

    # -------------------------------------------------------------------------
    # Sheet 2: Month-Wise Purchase
    # -------------------------------------------------------------------------
    row_idx = 1
    ws_months.cell(row=row_idx, column=1, value="Month-Wise Consolidated Purchase (Million INR)").font = section_font
    row_idx += 1
    merged_data = {}
    months_set = set()
    for row in data:
        sup = row.get("supplier_name") or row.get("supplier") or "-"
        prod = row.get("item_name") or row.get("item_code") or "-"
        amt = flt(row.get("base_amount") or 0)
        if row.get("is_return"): amt = -amt
        date_str = str(row.get("invoice_date") or "")
        try:
            d = frappe.utils.getdate(date_str)
            m_key = d.strftime("%b %Y"); m_sort = d.strftime("%Y%m")
        except:
            m_key = "Unknown"; m_sort = "000000"
        months_set.add((m_sort, m_key))
        key = f"{sup}|{prod}"
        if key not in merged_data: merged_data[key] = {"sup": sup, "prod": prod, "months": {}, "total": 0}
        merged_data[key]["months"][m_key] = merged_data[key]["months"].get(m_key, 0) + amt
        merged_data[key]["total"] += amt
    sorted_months = [x[1] for x in sorted(list(months_set), key=lambda x: x[0])]
    headers = ["S.No.", "Supplier", "Product"] + sorted_months + ["Total (M)"]
    for idx, h in enumerate(headers, start=1):
        cell = ws_months.cell(row=row_idx, column=idx, value=h); cell.font = header_font; cell.fill = header_fill; cell.alignment = Alignment(horizontal="center"); cell.border = table_border
    row_idx += 1
    
    total_months_sum = {m: 0 for m in sorted_months}
    grand_total_overall = 0
    for r_idx, row in enumerate(sorted(merged_data.values(), key=lambda x: x["total"], reverse=True)):
        row_fill = zebra_fill if r_idx % 2 != 0 else None
        c_no = ws_months.cell(row=row_idx, column=1, value=r_idx + 1); c1 = ws_months.cell(row=row_idx, column=2, value=row["sup"]); c2 = ws_months.cell(row=row_idx, column=3, value=row["prod"])
        for c in [c_no, c1, c2]: c.border = table_border; (row_fill and setattr(c, 'fill', row_fill))
        col_idx = 4
        for m_key in sorted_months:
            v_m = flt(row["months"].get(m_key, 0)) / 1000000
            total_months_sum[m_key] += flt(row["months"].get(m_key, 0))
            c = ws_months.cell(row=row_idx, column=col_idx, value=v_m); c.number_format = '"₹ "#,##0.00" M"'; c.border = table_border; c.alignment = Alignment(horizontal="right"); (row_fill and setattr(c, 'fill', row_fill))
            col_idx += 1
        tot_m = flt(row["total"]) / 1000000
        grand_total_overall += flt(row["total"])
        c_tot = ws_months.cell(row=row_idx, column=col_idx, value=tot_m); c_tot.number_format = '"₹ "#,##0.00" M"'; c_tot.font = Font(bold=True); c_tot.fill = PatternFill(start_color="ecf0f1", fill_type="solid"); c_tot.border = table_border; c_tot.alignment = Alignment(horizontal="right")
        row_idx += 1

    # Add Grand Total Row for Sheet 2
    ws_months.cell(row=row_idx, column=1, value="Grand Total").font = Font(bold=True)
    ws_months.merge_cells(start_row=row_idx, start_column=1, end_row=row_idx, end_column=3)
    for c_i in range(1, col_idx + 1): 
        cell = ws_months.cell(row=row_idx, column=c_i)
        cell.fill = PatternFill(start_color="f1c40f", fill_type="solid")
        cell.border = table_border
        cell.font = Font(bold=True)
        if c_i >= 4:
            cell.alignment = Alignment(horizontal="right")
    
    col_idx_tmp = 4
    for m_key in sorted_months:
        v_m = total_months_sum[m_key] / 1000000
        c = ws_months.cell(row=row_idx, column=col_idx_tmp, value=v_m)
        c.number_format = '"₹ "#,##0.00" M"'
        col_idx_tmp += 1
    c_all = ws_months.cell(row=row_idx, column=col_idx_tmp, value=grand_total_overall / 1000000)
    c_all.number_format = '"₹ "#,##0.00" M"'
    row_idx += 2

    # -------------------------------------------------------------------------
    # Sheet 3: Purchase Invoices List
    # -------------------------------------------------------------------------
    row_idx = 1
    ws_list.cell(row=row_idx, column=1, value="Detailed Purchase Invoices List (Million INR)").font = section_font
    row_idx += 1
    ui_columns = [
        {"label": "S.No.", "fieldname": "sr_no_idx", "width": 8},
        {"label": "Invoice ID", "fieldname": "invoice_id", "width": 18},
        {"label": "Date", "fieldname": "invoice_date", "width": 14},
        {"label": "Status", "fieldname": "status", "width": 14},
        {"label": "Supplier", "fieldname": "supplier_name", "width": 25},
        {"label": "Item", "fieldname": "item_name", "width": 20},
        {"label": "Qty", "fieldname": "qty", "width": 10},
        {"label": "Amount (M)", "fieldname": "base_amount", "width": 18},
    ]
    for idx, col in enumerate(ui_columns, start=1):
        cell = ws_list.cell(row=row_idx, column=idx, value=col["label"]); cell.font = header_font; cell.fill = header_fill; cell.alignment = Alignment(horizontal="center"); cell.border = table_border; ws_list.column_dimensions[get_column_letter(idx)].width = col["width"]
    row_idx += 1
    total_qty_list = 0
    total_amt_list = 0
    for r_idx, row in enumerate(data):
        row_fill = zebra_fill if r_idx % 2 != 0 else None
        for idx, col in enumerate(ui_columns, start=1):
            fname = col["fieldname"]; val = row.get(fname)
            if fname == "sr_no_idx": val = r_idx + 1
            cell = ws_list.cell(row=row_idx, column=idx); cell.border = table_border; (row_fill and setattr(cell, 'fill', row_fill))
            if fname in ["qty", "base_amount"]:
                num_val = flt(val or 0)
                if row.get("is_return") and fname == "base_amount": num_val = -num_val
                if fname == "base_amount":
                    num_val = num_val / 1000000; cell.number_format = '"₹ "#,##0.00" M"'
                    total_amt_list += (flt(val) if not row.get("is_return") else -flt(val))
                else: 
                    cell.number_format = "#,##0.00"
                    total_qty_list += flt(val)
                cell.value = num_val; cell.alignment = Alignment(horizontal="right")
            else:
                cell.value = str(val) if val is not None else ""; cell.alignment = Alignment(horizontal="left")
        row_idx += 1

    # Add Total Row for Sheet 3
    ws_list.cell(row=row_idx, column=1, value="Total").font = Font(bold=True)
    ws_list.merge_cells(start_row=row_idx, start_column=1, end_row=row_idx, end_column=6)
    for c_i in range(1, 9): 
        cell = ws_list.cell(row=row_idx, column=c_i)
        cell.fill = PatternFill(start_color="f1c40f", fill_type="solid")
        cell.border = table_border
        cell.font = Font(bold=True)
        if c_i >= 7:
            cell.alignment = Alignment(horizontal="right")
    
    c_q = ws_list.cell(row=row_idx, column=7, value=total_qty_list)
    c_q.number_format = "#,##0.00"
    
    c_a = ws_list.cell(row=row_idx, column=8, value=total_amt_list / 1000000)
    c_a.number_format = '"₹ "#,##0.00" M"'
    row_idx += 1
    for i in range(1, 10): ws_overview.column_dimensions[get_column_letter(i)].width = 20; ws_months.column_dimensions[get_column_letter(i)].width = 20
    output = BytesIO(); wb.save(output); output.seek(0)
    return {"filename": f"Purchase_Invoice_Dashboard_{frappe.utils.nowdate()}.xlsx", "filecontent": base64.b64encode(output.read()).decode()}
