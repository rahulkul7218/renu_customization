import frappe
from frappe import _
from frappe.utils import flt, getdate, nowdate, add_days, formatdate
import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
from io import BytesIO
import base64

@frappe.whitelist()
def export_to_pdf(html=None, orientation="Landscape"):
    if not html:
        frappe.throw(_("PDF content is empty"))
    
    options = {
        "page-size": "A4",
        "orientation": orientation,
        "margin-top": "10mm",
        "margin-right": "10mm",
        "margin-bottom": "10mm",
        "margin-left": "10mm",
        "encoding": "UTF-8",
        "no-outline": None,
    }
    
    pdf_content = frappe.utils.pdf.get_pdf(html, options)
    
    return {
        "filename": f"Supplier_Performance_{nowdate()}.pdf",
        "filecontent": base64.b64encode(pdf_content).decode(),
    }

@frappe.whitelist()
def get_po_items(purchase_order):
    if not purchase_order:
        return []
    
    if not frappe.has_permission("Purchase Order", "read"):
        frappe.throw(_("Not permitted to read Purchase Order"), frappe.PermissionError)
        
    items = frappe.get_all(
        "Purchase Order Item",
        filters={"parent": purchase_order},
        fields=["item_code", "item_name", "qty", "rate", "amount", "received_qty", "billed_amt", "description"],
        order_by="idx"
    )
    for item in items:
        item["billed_qty"] = flt(item.get("billed_amt")) / flt(item.get("rate")) if flt(item.get("rate")) else 0.0
    return items

def is_delivery_completed(row):
    qty = flt(row.get("qty"))
    received = flt(row.get("received_qty"))
    return (
        (qty > 0 and received >= qty)
        or flt(row.get("per_delivered")) >= 100
        or row.get("status") in ["Completed", "Closed"]
    )


def _parse_group_concat_dates(dates_str):
    if not dates_str:
        return []
    return [getdate(d.strip()) for d in str(dates_str).split(",") if d and str(d).strip()]


def get_actual_delivery_dates(row):
    dates = row.get("actual_delivery_dates") or []
    if isinstance(dates, str):
        dates = _parse_group_concat_dates(dates)
    elif dates:
        dates = [getdate(d) for d in dates if d]
    if not dates and row.get("actual_delivery_time"):
        dates = [getdate(row["actual_delivery_time"])]
    return sorted(dates)


def format_actual_delivery_display(row):
    dates = get_actual_delivery_dates(row)
    if not dates:
        return None
    return ", ".join(formatdate(d) for d in dates)


def _fetch_actual_delivery_map(po_item_field):
    """Return {po_item_name: [date, ...]} from all submitted purchase receipts."""
    actual_delivery_map = {}
    pr_data = frappe.db.sql(
        f"""
        SELECT
            pri.{po_item_field} as po_item_name,
            GROUP_CONCAT(DISTINCT pr.posting_date ORDER BY pr.posting_date SEPARATOR ',') as actual_delivery_dates
        FROM
            `tabPurchase Receipt` pr
        INNER JOIN
            `tabPurchase Receipt Item` pri ON pri.parent = pr.name
        WHERE
            pr.docstatus = 1
            AND pri.{po_item_field} IS NOT NULL
            AND pri.{po_item_field} != ''
        GROUP BY
            pri.{po_item_field}
        """,
        as_dict=True,
    )
    for r in pr_data:
        actual_delivery_map[r.po_item_name] = _parse_group_concat_dates(r.actual_delivery_dates)
    return actual_delivery_map


EXCLUDED_PO_STATUSES = ("Cancelled", "Draft")


def prepare_filters(filters):
    if not filters:
        filters = {}
    elif isinstance(filters, str):
        filters = frappe.parse_json(filters)
    
    if filters.get("date_range"):
        dr = filters.get("date_range")
        if isinstance(dr, list) and len(dr) == 2:
            filters["from_date"] = dr[0]
            filters["to_date"] = dr[1]

    if filters.get("fiscal_year") and (not filters.get("from_date") or not filters.get("to_date")):
        fy = frappe.get_doc("Fiscal Year", filters.get("fiscal_year"))
        if fy:
            filters["from_date"] = filters.get("from_date") or fy.year_start_date
            filters["to_date"] = filters.get("to_date") or fy.year_end_date

    return frappe._dict(filters)

@frappe.whitelist()
def get_dashboard_data(filters=None):
    if not frappe.has_permission("Purchase Order", "read"):
        frappe.throw(_("Not permitted to read Purchase Order"), frappe.PermissionError)

    filters = prepare_filters(filters)

    # Build Purchase Order filters using Frappe ORM
    po_filters = {"docstatus": 1}

    if filters.get("company"):
        po_filters["company"] = filters.get("company")

    if filters.get("from_date"):
        po_filters["transaction_date"] = [">=", filters.get("from_date")]

    if filters.get("to_date"):
        if "transaction_date" in po_filters:
            po_filters["transaction_date"] = ["between", [filters.get("from_date"), filters.get("to_date")]]
        else:
            po_filters["transaction_date"] = ["<=", filters.get("to_date")]

    if filters.get("purchase_order"):
        po_filters["name"] = filters.get("purchase_order")

    if filters.get("supplier"):
        po_filters["supplier"] = filters.get("supplier")
    elif filters.get("supplier_group"):
        suppliers = frappe.get_all("Supplier", filters={"supplier_group": filters.get("supplier_group")}, pluck="name")
        if suppliers:
            po_filters["supplier"] = ["in", suppliers]
        else:
            po_filters["supplier"] = ["in", [""]]

    if filters.get("open_po_details"):
        po_filters["status"] = ["in", ["To Receive and Bill", "To Receive", "To Bill"]]
    elif filters.get("status"):
        if filters.get("status") in EXCLUDED_PO_STATUSES:
            return {"summary": [], "results": []}
        po_filters["status"] = filters.get("status")
    else:
        po_filters["status"] = ["not in", list(EXCLUDED_PO_STATUSES)]

    # Fetch Purchase Orders
    purchase_orders = frappe.get_all(
        "Purchase Order",
        filters=po_filters,
        fields=["name", "supplier", "transaction_date", "status", "company"],
        order_by="transaction_date desc",
        limit_page_length=0
    )

    if not purchase_orders:
        return {"summary": [], "results": []}

    po_names = [po.name for po in purchase_orders]
    po_map = {po.name: po for po in purchase_orders}

    # Fetch Purchase Order Items
    poi_filters = {"parent": ["in", po_names]}

    if filters.get("expected_delivery_date"):
        poi_filters["schedule_date"] = filters.get("expected_delivery_date")

    po_items = frappe.get_all(
        "Purchase Order Item",
        filters=poi_filters,
        fields=["name", "parent", "item_code", "item_name", "qty", "rate", "amount",
                "received_qty", "billed_amt", "schedule_date"],
        order_by="parent, idx",
        limit_page_length=0
    )

    if not po_items:
        return {"summary": [], "results": []}

    # All actual delivery dates from Purchase Receipt (comma-separated per PO item)
    actual_delivery_map = {}
    try:
        actual_delivery_map = _fetch_actual_delivery_map("purchase_order_item")
    except Exception:
        try:
            actual_delivery_map = _fetch_actual_delivery_map("po_detail")
        except Exception:
            pass

    # Build combined results
    report_data = []
    for item in po_items:
        po = po_map.get(item.parent)
        if not po or po.status in EXCLUDED_PO_STATUSES:
            continue
        
        rate = flt(item.rate)
        qty = flt(item.qty)
        received_qty = flt(item.received_qty)
        billed_amt = flt(item.billed_amt)
        billed_qty = billed_amt / rate if rate > 0 else 0.0
        
        per_delivered = (received_qty / qty) * 100 if qty > 0 else 0.0
        per_billed = (billed_qty / qty) * 100 if qty > 0 else 0.0
        pending_qty = max(0, qty - received_qty)
        
        delivery_dates = actual_delivery_map.get(item.name, [])
        report_data.append({
            "name": po.name,
            "supplier": po.supplier,
            "transaction_date": po.transaction_date,
            "status": po.status,
            "po_item_name": item.name,
            "item_code": item.item_code,
            "item_name": item.item_name,
            "qty": qty,
            "rate": rate,
            "net_total": item.amount,
            "received_qty": received_qty,
            "pending_qty": pending_qty,
            "billed_qty": billed_qty,
            "per_delivered": per_delivered,
            "per_billed": per_billed,
            "schedule_date": item.schedule_date,
            "actual_delivery_dates": [str(d) for d in delivery_dates],
            "actual_delivery_time": str(delivery_dates[-1]) if delivery_dates else None,
        })

    # Filter by actual_delivery_time if specified
    if filters.get("actual_delivery_time"):
        adt = getdate(filters.get("actual_delivery_time"))
        report_data = [r for r in report_data if adt in get_actual_delivery_dates(r)]

    if not report_data:
        return {"summary": [], "results": []}

    results = []
    total_amount = 0
    total_overdue = 0
    total_due_next_15_days = 0
    unique_pos = set()

    today = getdate(nowdate())
    next_15_days = add_days(today, 15)

    for row in report_data:
        is_overdue = False
        due_next_15_days_flag = False

        row["due_days"] = "-"
        if row.get("schedule_date") and row.get("status") not in ["Completed", "Closed", "Cancelled"]:
            po_date = getdate(row.get("schedule_date"))
            pending_receipt = flt(row.get("received_qty")) < flt(row.get("qty"))

            if po_date < today and pending_receipt:
                is_overdue = True
                row["due_days"] = (today - po_date).days
            elif today <= po_date <= next_15_days:
                # Match "Orders Due in Next 15 Days" table: all open lines in window
                due_next_15_days_flag = True
                row["due_days"] = (po_date - today).days
            elif po_date > today and pending_receipt:
                row["due_days"] = (po_date - today).days

        row["is_overdue"] = is_overdue
        row["is_due_next_15_days"] = due_next_15_days_flag

        if filters.get("is_overdue") and not is_overdue:
            continue

        if filters.get("due_next_15_days") and not due_next_15_days_flag:
            continue

        unique_pos.add(row.get("name"))
        total_amount += flt(row.get("net_total"))
        if is_overdue:
            total_overdue += flt(row.get("net_total"))
        if due_next_15_days_flag:
            total_due_next_15_days += flt(row.get("net_total"))

        results.append(row)

    results.sort(key=lambda x: getdate(x.get("schedule_date")) if x.get("schedule_date") else today, reverse=True)

    summary = [
        {"label": _("Total Orders"), "value": len(unique_pos), "indicator": "blue", "fieldtype": "Int"},
        {"label": _("Total Net Amount"), "value": total_amount, "indicator": "green", "fieldtype": "Currency"},
        {"label": _("Overdue Amount"), "value": total_overdue, "indicator": "red", "fieldtype": "Currency"},
        {"label": _("Due Next 15 Days"), "value": total_due_next_15_days, "indicator": "orange", "fieldtype": "Currency"}
    ]

    supplier_totals = {}
    status_counts = {}

    for row in results:
        supp = row.get("supplier") or "Unknown"
        stat = row.get("status") or "Unknown"

        supplier_totals[supp] = supplier_totals.get(supp, 0) + flt(row.get("net_total"))
        status_counts[stat] = status_counts.get(stat, 0) + flt(row.get("net_total"))

    top_10_suppliers = sorted(supplier_totals.items(), key=lambda x: x[1], reverse=True)[:10]

    charts = {
        "top_10_suppliers": {
            "title": _("Top 10 Suppliers "),
            "data": {
                "labels": [x[0] for x in top_10_suppliers],
                "datasets": [{"name": "Amount", "values": [x[1] for x in top_10_suppliers]}]
            },
            "type": "donut",
            "colors": ["#3b82f6", "#8b5cf6", "#10b981", "#f59e0b", "#06b6d4", "#ef4444", "#6366f1", "#ec4899", "#84cc16", "#f97316"],
            "is_currency": True
        },
        "order_status": {
            "title": _("Order Status Wise Amount "),
            "data": {
                "labels": list(status_counts.keys()),
                "datasets": [{"name": "Amount", "values": list(status_counts.values())}]
            },
            "type": "donut",
            "colors": ["#10b981", "#f59e0b", "#3b82f6", "#ef4444", "#8b5cf6", "#6366f1", "#ec4899", "#84cc16"],
            "is_currency": True
        }
    }

    # Monthly Calculations
    months_set = set()
    monthly_lifecycle = {
        "Booked": {},
        "Delivered": {},
        "Pending": {},
        "Overdue": {}
    }
    month_wise_supplier = {}

    for row in results:
        try:
            d = getdate(row.get("transaction_date"))
            m_key = d.strftime("%b %Y")
            m_sort = d.strftime("%Y%m")
            months_set.add((m_sort, m_key))
        except:
            m_key = "Unknown"
        
        amt = flt(row.get("net_total"))
        supp = row.get("supplier") or "Unknown"

        monthly_lifecycle["Booked"][m_key] = monthly_lifecycle["Booked"].get(m_key, 0) + amt
        if row.get("status") in ["Completed", "Closed"]:
            monthly_lifecycle["Delivered"][m_key] = monthly_lifecycle["Delivered"].get(m_key, 0) + amt
        
        if row.get("status") not in ["Completed", "Closed", "Cancelled"]:
            monthly_lifecycle["Pending"][m_key] = monthly_lifecycle["Pending"].get(m_key, 0) + amt
            if row.get("is_overdue"):
                monthly_lifecycle["Overdue"][m_key] = monthly_lifecycle["Overdue"].get(m_key, 0) + amt

        if supp not in month_wise_supplier:
            month_wise_supplier[supp] = {"supplier": supp, "months": {}, "total": 0}
        month_wise_supplier[supp]["months"][m_key] = month_wise_supplier[supp]["months"].get(m_key, 0) + amt
        month_wise_supplier[supp]["total"] += amt

    due_next_15_days = [row for row in results if row.get("is_due_next_15_days")]

    sorted_months = [{"key": x[1], "sort": x[0]} for x in sorted(list(months_set), key=lambda x: x[0])]

    return {
        "summary": summary,
        "results": results,
        "charts": charts,
        "months": sorted_months,
        "monthly_lifecycle": monthly_lifecycle,
        "month_wise_supplier": sorted(month_wise_supplier.values(), key=lambda x: x["total"], reverse=True),
        "due_next_15_days": sorted(due_next_15_days, key=lambda x: x.get("schedule_date") if x.get("schedule_date") else today)
    }

def _excel_styles():
    thin_side = Side(style="thin")
    return {
        "header_font": Font(bold=True, color="FFFFFF"),
        "header_fill": PatternFill(start_color="2c3e50", fill_type="solid"),
        "title_font": Font(bold=True, size=14),
        "section_font": Font(bold=True, size=12),
        "table_border": Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side),
        "zebra_fill": PatternFill(start_color="f8f9fa", fill_type="solid"),
    }


def _write_excel_headers(ws, headers, styles, row_idx=1):
    for idx, header in enumerate(headers, start=1):
        cell = ws.cell(row=row_idx, column=idx, value=header)
        cell.font = styles["header_font"]
        cell.fill = styles["header_fill"]
        cell.alignment = Alignment(horizontal="center")
        cell.border = styles["table_border"]


def _write_excel_row(ws, row_idx, values, styles, amount_col=None, date_cols=None, qty_cols=None, pct_cols=None):
    row_fill = styles["zebra_fill"] if row_idx % 2 == 0 else None
    date_cols = date_cols or set()
    qty_cols = qty_cols or set()
    pct_cols = pct_cols or set()
    for c_idx, val in enumerate(values, start=1):
        cell = ws.cell(row=row_idx, column=c_idx, value=val)
        cell.border = styles["table_border"]
        if row_fill:
            cell.fill = row_fill
        if c_idx in date_cols and val:
            cell.number_format = "DD-MM-YYYY"
        if c_idx in qty_cols:
            cell.number_format = "#,##0.##"
            cell.alignment = Alignment(horizontal="right")
        if c_idx in pct_cols:
            cell.number_format = "0"
            cell.alignment = Alignment(horizontal="center")
        if amount_col and c_idx == amount_col:
            cell.number_format = "#,##0.00"
            cell.alignment = Alignment(horizontal="right")


def _total_row_style(styles):
    return {
        "font": Font(bold=True, size=11),
        "fill": PatternFill(start_color="e2e8f0", fill_type="solid"),
        "border": styles["table_border"],
    }


def _write_excel_total_row(ws, row_idx, label, amount_col, amount_value, styles):
    """Total label merged across columns 1..amount_col-1; amount in last column."""
    total_style = _total_row_style(styles)
    label_end_col = max(1, amount_col - 1)

    for col in range(1, amount_col + 1):
        cell = ws.cell(row=row_idx, column=col)
        cell.fill = total_style["fill"]
        cell.border = total_style["border"]

    if label_end_col > 1:
        ws.merge_cells(start_row=row_idx, start_column=1, end_row=row_idx, end_column=label_end_col)

    label_cell = ws.cell(row=row_idx, column=1, value=label)
    label_cell.font = total_style["font"]
    label_cell.fill = total_style["fill"]
    label_cell.border = total_style["border"]
    label_cell.alignment = Alignment(horizontal="right", vertical="center")

    amount_cell = ws.cell(row=row_idx, column=amount_col, value=amount_value)
    amount_cell.font = total_style["font"]
    amount_cell.fill = total_style["fill"]
    amount_cell.border = total_style["border"]
    amount_cell.number_format = "#,##0.00"
    amount_cell.alignment = Alignment(horizontal="right")


def _write_chart_tables_on_overview(ws, dashboard_data, styles, start_row):
    charts = dashboard_data.get("charts") or {}
    row_idx = start_row
    chart_headers = ["Label", "Amount (M)", "% Share"]

    for chart_obj in charts.values():
        title = chart_obj.get("title") or "Chart"
        labels = chart_obj.get("data", {}).get("labels") or []
        values = (chart_obj.get("data", {}).get("datasets") or [{}])[0].get("values") or []
        is_currency = chart_obj.get("is_currency")
        chart_total = sum(flt(v) for v in values)

        ws.cell(row=row_idx, column=1, value=title).font = styles["section_font"]
        row_idx += 1
        _write_excel_headers(ws, chart_headers, styles, row_idx=row_idx)
        row_idx += 1

        for label, val in zip(labels, values):
            amount = flt(val) / 1000000 if is_currency else flt(val)
            pct = (flt(val) / chart_total * 100) if chart_total else 0
            _write_excel_row(
                ws, row_idx, [label, amount, round(pct, 1)],
                styles, amount_col=2, pct_cols={3},
            )
            row_idx += 1

        _write_excel_row(
            ws, row_idx, ["TOTAL", chart_total / 1000000 if is_currency else chart_total, 100],
            styles, amount_col=2, pct_cols={3},
        )
        for col in range(1, 4):
            cell = ws.cell(row=row_idx, column=col)
            cell.font = Font(bold=True)
            cell.fill = PatternFill(start_color="e2e8f0", fill_type="solid")
        row_idx += 2

    return row_idx


def _autofit_columns(ws, min_widths=None):
    """Set column width from content, respecting optional minimum widths per column letter."""
    min_widths = min_widths or {}
    for col in ws.columns:
        max_length = 0
        column = col[0].column_letter
        for cell in col:
            try:
                if cell.value is not None:
                    max_length = max(max_length, len(str(cell.value)))
            except Exception:
                pass
        width = max(min_widths.get(column, 0), max_length + 3)
        ws.column_dimensions[column].width = min(width, 55)


def _apply_header_column_widths(ws, headers, width_by_header):
    for idx, header in enumerate(headers, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = width_by_header.get(header, 14)


def _min_widths_for_headers(headers, width_by_header):
    return {
        get_column_letter(idx): width_by_header.get(header, 14)
        for idx, header in enumerate(headers, start=1)
    }


MONTH_SHEET_WIDTHS = {
    "S.No.": 8,
    "Supplier": 32,
    "Total (M)": 14,
}

ORDER_SHEET_WIDTHS = {
    "S.No.": 8,
    "PO No": 22,
    "Supplier": 28,
    "Item Code": 16,
    "Item Name": 42,
    "Order Date": 14,
    "Expected Delivery": 16,
    "Actual Delivery": 22,
    "Days Left": 12,
    "Status": 22,
    "% Received": 12,
    "% Billed": 12,
    "Order Qty": 12,
    "Received Qty": 14,
    "Pending Qty": 14,
    "Net Total (M)": 14,
}


def _write_section_title(ws, row_idx, title, styles):
    ws.cell(row=row_idx, column=1, value=title).font = styles["section_font"]
    return row_idx + 1


def _write_overview_sheet(ws, dashboard_data, styles, results=None, due_rows=None):
    summary = dashboard_data.get("summary") or []
    ws.cell(row=1, column=1, value="Supplier Performance Dashboard Overview").font = styles["title_font"]
    ws.cell(row=1, column=4, value="Generated On: " + str(nowdate()))
    
    row_idx = 3
    row_idx = _write_section_title(ws, row_idx, "Operational Summary (Million INR)", styles)
    
    colors = {"blue": "3498db", "green": "2ecc71", "red": "e74c3c", "orange": "e67e22"}
    for i, s in enumerate(summary):
        r = row_idx + (i // 4) * 3
        c = 1 + (i % 4) * 2
        bg_color = colors.get((s.get("indicator") or "blue").lower(), "3498db")
        
        cell_l = ws.cell(row=r, column=c, value=s.get("label"))
        cell_l.font = Font(bold=True, color="FFFFFF")
        cell_l.fill = PatternFill(start_color=bg_color, fill_type="solid")
        cell_l.alignment = Alignment(horizontal="center")
        
        val = flt(s.get("value"))
        if s.get("fieldtype") == "Currency":
            val = val / 1000000
        
        cell_v = ws.cell(row=r + 1, column=c, value=val)
        cell_v.font = Font(bold=True, size=12)
        if s.get("fieldtype") == "Currency":
            cell_v.number_format = '"₹ "#,##0.00" M"'
        cell_v.alignment = Alignment(horizontal="center")

        ws.merge_cells(start_row=r, start_column=c, end_row=r, end_column=c + 1)
        ws.merge_cells(start_row=r + 1, start_column=c, end_row=r + 1, end_column=c + 1)

    row_idx = row_idx + 4
    row_idx = _write_section_title(ws, row_idx, "Visual Analytics", styles)
    row_idx = _write_chart_tables_on_overview(ws, dashboard_data, styles, row_idx)

    if results is not None and due_rows is not None:
        row_idx += 1
        row_idx = _write_section_title(ws, row_idx, "Month-Wise Booking Breakdown", styles)
        row_idx = _write_month_table(ws, dashboard_data, styles, row_idx)

        row_idx += 1
        row_idx = _write_section_title(ws, row_idx, "Orders Due in Next 15 Days", styles)
        row_idx = _write_due_orders_table(ws, due_rows, styles, row_idx)

        row_idx += 1
        row_idx = _write_section_title(ws, row_idx, "Detailed Supplier Orders List ", styles)
        _write_detailed_orders_table(ws, results, styles, row_idx)

    _autofit_columns(ws)


def _write_month_table(ws, dashboard_data, styles, start_row=1):
    months = dashboard_data.get("months") or []
    headers = ["S.No.", "Supplier"] + [m["key"] for m in months] + ["Total (M)"]
    _write_excel_headers(ws, headers, styles, row_idx=start_row)

    month_widths = dict(MONTH_SHEET_WIDTHS)
    for m in months:
        month_widths[m["key"]] = 14

    if start_row == 1:
        _apply_header_column_widths(ws, headers, month_widths)

    amount_cols = set(range(3, len(headers)))
    month_rows = dashboard_data.get("month_wise_supplier") or []
    month_totals = {m["key"]: 0 for m in months}
    grand_total = 0
    data_start = start_row + 1
    for i, row in enumerate(month_rows):
        r_idx = data_start + i
        vals = [i + 1, row.get("supplier")]
        for m in months:
            amt = flt(row.get("months", {}).get(m["key"], 0))
            month_totals[m["key"]] += amt
            vals.append(amt / 1000000)
        row_total = flt(row.get("total"))
        grand_total += row_total
        vals.append(row_total / 1000000)
        _write_excel_row(ws, r_idx, vals, styles, amount_col=len(vals), qty_cols=amount_cols)

    total_row_idx = data_start + len(month_rows)
    ws.merge_cells(start_row=total_row_idx, start_column=1, end_row=total_row_idx, end_column=2)
    total_label = ws.cell(row=total_row_idx, column=1, value="GRAND TOTAL")
    total_style = _total_row_style(styles)
    total_label.font = total_style["font"]
    total_label.fill = total_style["fill"]
    total_label.border = total_style["border"]
    total_label.alignment = Alignment(horizontal="right")
    ws.cell(row=total_row_idx, column=2).fill = total_style["fill"]
    ws.cell(row=total_row_idx, column=2).border = total_style["border"]
    for col_idx, m in enumerate(months, start=3):
        cell = ws.cell(row=total_row_idx, column=col_idx, value=month_totals[m["key"]] / 1000000)
        cell.font = total_style["font"]
        cell.fill = total_style["fill"]
        cell.border = total_style["border"]
        cell.number_format = "#,##0.00"
        cell.alignment = Alignment(horizontal="right")
    grand_cell = ws.cell(row=total_row_idx, column=len(headers), value=grand_total / 1000000)
    grand_cell.font = total_style["font"]
    grand_cell.fill = total_style["fill"]
    grand_cell.border = total_style["border"]
    grand_cell.number_format = "#,##0.00"
    grand_cell.alignment = Alignment(horizontal="right")

    if start_row == 1:
        _autofit_columns(ws, min_widths=_min_widths_for_headers(headers, month_widths))
    return total_row_idx + 1


def _write_month_sheet(ws, dashboard_data, styles):
    _write_month_table(ws, dashboard_data, styles, start_row=1)


def _as_excel_date(val):
    if not val:
        return None
    try:
        return getdate(val)
    except Exception:
        return val


def _order_row_values(row, serial_no, include_days_left=False):
    pending_qty = flt(row.get("pending_qty"))
    if not pending_qty and row.get("qty") is not None:
        pending_qty = max(0, flt(row.get("qty")) - flt(row.get("received_qty")))

    vals = [
        serial_no,
        row.get("name"),
        row.get("supplier"),
        row.get("item_code"),
        row.get("item_name"),
        _as_excel_date(row.get("transaction_date")),
        _as_excel_date(row.get("schedule_date")),
        format_actual_delivery_display(row),
    ]
    if include_days_left:
        due_days = row.get("due_days")
        vals.append(f"{due_days} Days" if due_days != "-" else due_days)
    vals.extend([
        row.get("status"),
        int(row.get("per_delivered") or 0),
        int(row.get("per_billed") or 0),
        flt(row.get("qty")),
        flt(row.get("received_qty")),
        pending_qty,
        flt(row.get("net_total")) / 1000000,
    ])
    return vals


DETAILED_ORDER_HEADERS = [
    "S.No.", "PO No", "Supplier", "Item Code", "Item Name", "Order Date",
    "Expected Delivery", "Actual Delivery", "Status", "% Received", "% Billed",
    "Order Qty", "Received Qty", "Pending Qty", "Net Total (M)",
]

DUE_ORDER_HEADERS = [
    "S.No.", "PO No", "Supplier", "Item Code", "Item Name", "Order Date",
    "Expected Delivery", "Actual Delivery", "Days Left", "Status", "% Received", "% Billed",
    "Order Qty", "Received Qty", "Pending Qty", "Net Total (M)",
]


def _order_sheet_col_sets(headers):
    """Return column index sets for date, qty, pct formatting based on header positions."""
    date_headers = {"Order Date", "Expected Delivery"}
    qty_headers = {"Order Qty", "Received Qty", "Pending Qty"}
    pct_headers = {"% Received", "% Billed"}
    date_cols, qty_cols, pct_cols = set(), set(), set()
    for idx, header in enumerate(headers, start=1):
        if header in date_headers:
            date_cols.add(idx)
        if header in qty_headers:
            qty_cols.add(idx)
        if header in pct_headers:
            pct_cols.add(idx)
    return date_cols, qty_cols, pct_cols


def _write_detailed_orders_table(ws, rows, styles, start_row=1):
    _write_excel_headers(ws, DETAILED_ORDER_HEADERS, styles, row_idx=start_row)
    if start_row == 1:
        _apply_header_column_widths(ws, DETAILED_ORDER_HEADERS, ORDER_SHEET_WIDTHS)
    date_cols, qty_cols, pct_cols = _order_sheet_col_sets(DETAILED_ORDER_HEADERS)
    data_rows = rows or []
    data_start = start_row + 1
    for i, row in enumerate(data_rows):
        r_idx = data_start + i
        _write_excel_row(
            ws, r_idx, _order_row_values(row, i + 1), styles,
            amount_col=len(DETAILED_ORDER_HEADERS),
            date_cols=date_cols, qty_cols=qty_cols, pct_cols=pct_cols,
        )
    total_amount = sum(flt(r.get("net_total")) for r in data_rows) / 1000000
    _write_excel_total_row(
        ws, data_start + len(data_rows), "TOTAL BOOKED VALUE", len(DETAILED_ORDER_HEADERS), total_amount, styles,
    )
    if start_row == 1:
        _autofit_columns(ws, min_widths=_min_widths_for_headers(DETAILED_ORDER_HEADERS, ORDER_SHEET_WIDTHS))
    return data_start + len(data_rows) + 1


def _write_detailed_orders_sheet(ws, rows, styles):
    _write_detailed_orders_table(ws, rows, styles, start_row=1)


def _write_due_orders_table(ws, rows, styles, start_row=1, total_label="TOTAL DUE VALUE"):
    _write_excel_headers(ws, DUE_ORDER_HEADERS, styles, row_idx=start_row)
    if start_row == 1:
        _apply_header_column_widths(ws, DUE_ORDER_HEADERS, ORDER_SHEET_WIDTHS)
    date_cols, qty_cols, pct_cols = _order_sheet_col_sets(DUE_ORDER_HEADERS)
    data_rows = rows or []
    data_start = start_row + 1
    for i, row in enumerate(data_rows):
        r_idx = data_start + i
        _write_excel_row(
            ws, r_idx, _order_row_values(row, i + 1, include_days_left=True), styles,
            amount_col=len(DUE_ORDER_HEADERS),
            date_cols=date_cols, qty_cols=qty_cols, pct_cols=pct_cols,
        )
    total_amount = sum(flt(r.get("net_total")) for r in data_rows) / 1000000
    _write_excel_total_row(
        ws, data_start + len(data_rows), total_label, len(DUE_ORDER_HEADERS), total_amount, styles,
    )
    if start_row == 1:
        _autofit_columns(ws, min_widths=_min_widths_for_headers(DUE_ORDER_HEADERS, ORDER_SHEET_WIDTHS))
    return data_start + len(data_rows) + 1


def _write_due_orders_sheet(ws, rows, styles):
    _write_due_orders_table(ws, rows, styles, start_row=1)


def _write_overdue_orders_sheet(ws, rows, styles):
    _write_due_orders_table(ws, rows, styles, start_row=1, total_label="TOTAL OVERDUE VALUE")


@frappe.whitelist()
def export_to_excel(filters=None, export_type="all"):
    filters = prepare_filters(filters)
    dashboard_data = get_dashboard_data(filters)
    results = dashboard_data.get("results") or []
    due_rows = dashboard_data.get("due_next_15_days") or []
    overdue_rows = [r for r in results if r.get("is_overdue")]

    if not results and not dashboard_data.get("month_wise_supplier"):
        frappe.throw(_("No data to export"))

    styles = _excel_styles()
    wb = openpyxl.Workbook()

    if export_type == "all":
        ws_overview = wb.active
        ws_overview.title = "Overview"
        _write_overview_sheet(ws_overview, dashboard_data, styles)

        ws_months = wb.create_sheet("Month-Wise Booking")
        _write_month_sheet(ws_months, dashboard_data, styles)

        ws_due = wb.create_sheet("Due in 15 Days")
        _write_due_orders_sheet(ws_due, due_rows, styles)

        ws_overdue = wb.create_sheet("Overdue Orders")
        _write_overdue_orders_sheet(ws_overdue, overdue_rows, styles)

        ws_list = wb.create_sheet("Detailed Orders")
        _write_detailed_orders_sheet(ws_list, results, styles)
    elif export_type == "summary":
        ws_months = wb.active
        ws_months.title = "Month-Wise Booking"
        _write_month_sheet(ws_months, dashboard_data, styles)
    elif export_type == "due":
        ws_due = wb.active
        ws_due.title = "Due in 15 Days"
        _write_due_orders_sheet(ws_due, due_rows, styles)
    elif export_type == "overdue":
        ws_overdue = wb.active
        ws_overdue.title = "Overdue Orders"
        _write_overdue_orders_sheet(ws_overdue, overdue_rows, styles)
    elif export_type == "detail":
        ws_list = wb.active
        ws_list.title = "Detailed Orders"
        _write_detailed_orders_sheet(ws_list, results, styles)
    else:
        frappe.throw(_("Invalid export type"))

    output = BytesIO()
    wb.save(output)
    output.seek(0)
    
    filenames = {
        "all": f"Supplier_Performance_{nowdate()}.xlsx",
        "summary": f"Supplier_Performance_Month_Wise_{nowdate()}.xlsx",
        "due": f"Supplier_Performance_Due_15_Days_{nowdate()}.xlsx",
        "overdue": f"Supplier_Performance_Overdue_{nowdate()}.xlsx",
        "detail": f"Supplier_Performance_Detailed_{nowdate()}.xlsx",
    }

    return {
        "filename": filenames.get(export_type, filenames["all"]),
        "filecontent": base64.b64encode(output.getvalue()).decode(),
    }
