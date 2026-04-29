import frappe
from frappe import _
from frappe.utils import flt, getdate, nowdate, add_days
from erpnext.buying.report.purchase_order_analysis.purchase_order_analysis import execute as po_analysis_execute
import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
from io import BytesIO
import base64

@frappe.whitelist()
def export_to_pdf(html):
    frappe.response.filename = "supplier_order_dashboard.pdf"
    frappe.response.type = "binary"
    frappe.response.filecontent = frappe.utils.pdf.get_pdf(html, {"orientation": "Landscape"})

def prepare_filters(filters):
    if not filters:
        filters = {}
    elif isinstance(filters, str):
        filters = frappe.parse_json(filters)
    return frappe._dict(filters)

@frappe.whitelist()
def get_dashboard_data(filters=None):
    filters = prepare_filters(filters)
    
    report_filters = frappe._dict({
        "group_by_po": 1
    })
    
    if filters.get("company"):
        report_filters["company"] = filters.get("company")
    
    if filters.get("from_date"):
        report_filters["from_date"] = filters.get("from_date")
    
    if filters.get("to_date"):
        report_filters["to_date"] = filters.get("to_date")
        
    if filters.get("purchase_order"):
        report_filters["name"] = [filters.get("purchase_order")]
        
    if filters.get("status"):
        report_filters["status"] = [filters.get("status")]
        
    try:
        columns, report_data, _msg, chart_data = po_analysis_execute(report_filters)
    except Exception as e:
        frappe.log_error(title="Supplier Dashboard Error", message=frappe.get_traceback())
        report_data = []
        
    if not report_data:
        return {"summary": [], "results": []}

    # Fetch custom fields for the fetched POs
    po_names = [d.get("purchase_order") for d in report_data if d.get("purchase_order")]
    
    po_details_map = {}
    if po_names:
        fields = ["name", "supplier_agreed_time", "delivery_time_as_per_po", "actual_delivery_time"]
        # Only fetch fields that exist to avoid errors
        meta = frappe.get_meta("Purchase Order")
        valid_fields = ["name"]
        for f in fields[1:]:
            if meta.has_field(f):
                valid_fields.append(f)
                
        po_details = frappe.get_all("Purchase Order", filters={"name": ("in", po_names)}, fields=valid_fields, limit_page_length=None)
        for po in po_details:
            po_details_map[po.name] = po

    results = []
    total_pos = 0
    total_amount = 0
    total_overdue = 0
    total_due_next_week = 0
    
    today = getdate(nowdate())
    next_week = add_days(today, 7)
    
    for row in report_data:
        po_name = row.get("purchase_order")
        po_detail = po_details_map.get(po_name, {})
        
        # Merge fields
        row["name"] = po_name
        row["transaction_date"] = row.get("date")
        row["schedule_date"] = row.get("required_date")
        row["net_total"] = row.get("amount") # Base amount
        
        row["supplier_agreed_time"] = po_detail.get("supplier_agreed_time")
        row["delivery_time_as_per_po"] = po_detail.get("delivery_time_as_per_po")
        row["actual_delivery_time"] = po_detail.get("actual_delivery_time")
        
        # Post-query filters
        if filters.get("supplier") and row.get("supplier") != filters.get("supplier"):
            continue
            
        if filters.get("expected_delivery_date") and str(row.get("schedule_date")) != str(filters.get("expected_delivery_date")):
            continue
            
        if filters.get("actual_delivery_time") and str(row.get("actual_delivery_time")) != str(filters.get("actual_delivery_time")):
            continue
            
        if filters.get("delivery_time_as_per_po") and str(row.get("delivery_time_as_per_po")) != str(filters.get("delivery_time_as_per_po")):
            continue
            
        if filters.get("supplier_agreed_time") and str(row.get("supplier_agreed_time")) != str(filters.get("supplier_agreed_time")):
            continue
            
        if filters.get("open_po_details") and row.get("status") not in ["Draft", "To Receive and Bill", "To Receive", "To Bill"]:
            continue

        is_overdue = False
        due_next_week_flag = False
        
        if row.get("schedule_date") and row.get("status") not in ["Completed", "Closed", "Cancelled"]:
            po_date = getdate(row.get("schedule_date"))
            if po_date < today:
                is_overdue = True
            elif today <= po_date <= next_week:
                due_next_week_flag = True
                
        row["is_overdue"] = is_overdue
        
        if filters.get("is_overdue") and not is_overdue:
            continue
            
        if filters.get("due_next_week") and not due_next_week_flag:
            continue
            
        total_pos += 1
        total_amount += flt(row.get("net_total"))
        if is_overdue:
            total_overdue += 1
        if due_next_week_flag:
            total_due_next_week += 1
            
        results.append(row)

    results.sort(key=lambda x: getdate(x.get("schedule_date")) if x.get("schedule_date") else today, reverse=True)

    summary = [
        {"label": _("Total Orders"), "value": total_pos, "indicator": "blue", "fieldtype": "Int"},
        {"label": _("Total Net Amount"), "value": total_amount, "indicator": "green", "fieldtype": "Currency"},
        {"label": _("Overdue Orders"), "value": total_overdue, "indicator": "red", "fieldtype": "Int"},
        {"label": _("Due in Next Week"), "value": total_due_next_week, "indicator": "orange", "fieldtype": "Int"}
    ]

    supplier_totals = {}
    status_counts = {}
    
    for row in results:
        supp = row.get("supplier") or "Unknown"
        stat = row.get("status") or "Unknown"
        
        supplier_totals[supp] = supplier_totals.get(supp, 0) + flt(row.get("net_total"))
        status_counts[stat] = status_counts.get(stat, 0) + 1

    top_10_suppliers = sorted(supplier_totals.items(), key=lambda x: x[1], reverse=True)[:10]
    
    charts = {
        "top_10_suppliers": {
            "data": {
                "labels": [x[0] for x in top_10_suppliers],
                "datasets": [{"name": "Amount", "values": [x[1] for x in top_10_suppliers]}]
            },
            "type": "bar",
            "colors": ["#3b82f6", "#8b5cf6", "#10b981", "#f59e0b", "#06b6d4", "#ef4444", "#6366f1", "#ec4899", "#84cc16", "#f97316"],
            "is_currency": True
        },
        "order_status": {
            "data": {
                "labels": list(status_counts.keys()),
                "datasets": [{"name": "Count", "values": list(status_counts.values())}]
            },
            "type": "donut",
            "colors": ["#10b981", "#f59e0b", "#3b82f6", "#ef4444", "#8b5cf6", "#6366f1", "#ec4899", "#84cc16"]
        }
    }

    return {
        "summary": summary,
        "results": results,
        "charts": charts
    }

@frappe.whitelist()
def export_to_excel(filters=None, export_type="all"):
    filters = prepare_filters(filters)
    dashboard_data = get_dashboard_data(filters)
    data = dashboard_data.get("results")
    
    if not data:
        return None

    wb = openpyxl.Workbook()
    
    # Conditionally create sheets based on export_type
    if export_type == "all":
        summary_ws = wb.active
        summary_ws.title = "Summary & Analytics"
        ws_months = wb.create_sheet("Month-Wise Booking")
        ws_list = wb.create_sheet("Supplier Orders List")
    elif export_type == "summary":
        ws_months = wb.active
        ws_months.title = "Month-Wise Booking"
        summary_ws = wb.create_sheet("Dummy1")
        ws_list = wb.create_sheet("Dummy2")
    elif export_type == "detail":
        ws_list = wb.active
        ws_list.title = "Supplier Orders List"
        summary_ws = wb.create_sheet("Dummy1")
        ws_months = wb.create_sheet("Dummy2")
    
    summary_ws.column_dimensions["A"].width = 30
    summary_ws.column_dimensions["B"].width = 20

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="2c3e50", fill_type="solid")
    section_font = Font(bold=True, size=12)
    thin_side = Side(style='thin')
    table_border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)
    zebra_fill = PatternFill(start_color="f8f9fa", fill_type="solid")

    if export_type == "all":
        r_idx = 1
        summary_ws.cell(row=r_idx, column=1, value="Supplier Order Dashboard - Summary").font = section_font
        r_idx += 2
        
        summary_ws.cell(row=r_idx, column=1, value="Key Performance Indicators").font = section_font
        r_idx += 2
        
        # Map colors to match dashboard indicators
        indicator_colors = {
            "blue": "3b82f6",
            "green": "10b981",
            "red": "ef4444",
            "orange": "f59e0b",
            "purple": "8b5cf6",
            "cyan": "06b6d4"
        }
        
        for i, item in enumerate(dashboard_data.get("summary", [])):
            # Calculate row and column for 2x2 grid
            grid_row = r_idx + (i // 2) * 3
            grid_col = 1 + (i % 2) * 2
            
            indicator = item.get("indicator", "blue").lower()
            bg_color = indicator_colors.get(indicator, "3b82f6")
            
            # Label cell
            cell_l = summary_ws.cell(row=grid_row, column=grid_col, value=item.get("label"))
            cell_l.font = Font(bold=True, color="FFFFFF")
            cell_l.fill = PatternFill(start_color=bg_color, fill_type="solid")
            cell_l.border = table_border
            cell_l.alignment = Alignment(horizontal="center")
            
            # Value cell
            val = item.get("value")
            val_col = grid_col + 1
            if item.get("fieldtype") == "Currency":
                val_display = flt(val) / 1000000
                val_cell = summary_ws.cell(row=grid_row, column=val_col, value=val_display)
                val_cell.number_format = '[$₹-en-IN] #,##0.00 "M"'
            else:
                val_cell = summary_ws.cell(row=grid_row, column=val_col, value=val)
                
            val_cell.font = Font(bold=True)
            val_cell.border = table_border
            val_cell.alignment = Alignment(horizontal="center")
            
            # Add a colored bottom border to the value cell to match dashboard card style
            val_cell.border = Border(
                left=thin_side, 
                right=thin_side, 
                top=thin_side, 
                bottom=Side(style='medium', color=bg_color)
            )

        # Adjust r_idx for next section (Status Breakdown)
        r_idx += 7
        
        summary_ws.column_dimensions["C"].width = 30
        summary_ws.column_dimensions["D"].width = 20
        summary_ws.cell(row=r_idx, column=1, value="Order Status Breakdown").font = Font(bold=True)
        r_idx += 1
        summary_ws.cell(row=r_idx, column=1, value="Status").font = header_font; summary_ws.cell(row=r_idx, column=1).fill = header_fill; summary_ws.cell(row=r_idx, column=1).border = table_border
        summary_ws.cell(row=r_idx, column=2, value="Count").font = header_font; summary_ws.cell(row=r_idx, column=2).fill = header_fill; summary_ws.cell(row=r_idx, column=2).border = table_border
        summary_ws.cell(row=r_idx, column=3, value="Share %").font = header_font; summary_ws.cell(row=r_idx, column=3).fill = header_fill; summary_ws.cell(row=r_idx, column=3).border = table_border
        r_idx += 1
        
        charts = dashboard_data.get("charts", {})
        if "order_status" in charts:
            labels = charts["order_status"]["data"]["labels"]
            vals = charts["order_status"]["data"]["datasets"][0]["values"]
            total_val = sum(vals) or 1
            for l, v in zip(labels, vals):
                summary_ws.cell(row=r_idx, column=1, value=l).border = table_border
                summary_ws.cell(row=r_idx, column=2, value=v).border = table_border
                share_c = summary_ws.cell(row=r_idx, column=3, value=(v / total_val))
                share_c.border = table_border; share_c.number_format = '0.0%'
                r_idx += 1

        r_idx += 2
        summary_ws.cell(row=r_idx, column=1, value="Top 10 Suppliers by Value").font = Font(bold=True)
        r_idx += 1
        summary_ws.cell(row=r_idx, column=1, value="Supplier").font = header_font; summary_ws.cell(row=r_idx, column=1).fill = header_fill; summary_ws.cell(row=r_idx, column=1).border = table_border
        summary_ws.cell(row=r_idx, column=2, value="Amount (M)").font = header_font; summary_ws.cell(row=r_idx, column=2).fill = header_fill; summary_ws.cell(row=r_idx, column=2).border = table_border
        summary_ws.cell(row=r_idx, column=3, value="Share %").font = header_font; summary_ws.cell(row=r_idx, column=3).fill = header_fill; summary_ws.cell(row=r_idx, column=3).border = table_border
        r_idx += 1
        
        if "top_10_suppliers" in charts:
            labels = charts["top_10_suppliers"]["data"]["labels"]
            vals = charts["top_10_suppliers"]["data"]["datasets"][0]["values"]
            total_val = sum(vals) or 1
            for l, v in zip(labels, vals):
                summary_ws.cell(row=r_idx, column=1, value=l).border = table_border
                val_c = summary_ws.cell(row=r_idx, column=2, value=v / 1000000)
                val_c.border = table_border
                val_c.number_format = '[$₹-en-IN] #,##0.00 "M"'
                share_c = summary_ws.cell(row=r_idx, column=3, value=(v / total_val))
                share_c.border = table_border; share_c.number_format = '0.0%'
                r_idx += 1

    if export_type in ["all", "summary"]:
        ws_months = wb.get_sheet_by_name("Month-Wise Booking") if "Month-Wise Booking" in wb.sheetnames else wb.create_sheet("Month-Wise Booking")
        row_idx = 1
        ws_months.cell(row=row_idx, column=1, value="Month-Wise Booking Breakdown (Million INR)").font = section_font
        row_idx += 2
        
        merged_data = {}
        months_set = set()
        for row in data:
            supp = row.get("supplier") or "-"
            amt = flt(row.get("net_total") or 0)
            try:
                d = frappe.utils.getdate(row.get("transaction_date"))
                m_key, m_sort = d.strftime("%b %Y"), d.strftime("%Y%m")
            except: m_key, m_sort = "Unknown", "000000"
            
            if row.get("transaction_date"):
                months_set.add((m_sort, m_key))
                
            if supp not in merged_data: merged_data[supp] = {"supp": supp, "months": {}, "total": 0}
            merged_data[supp]["months"][m_key] = merged_data[supp]["months"].get(m_key, 0) + amt
            merged_data[supp]["total"] += amt
            
        sorted_months = [x[1] for x in sorted(list(months_set), key=lambda x: x[0])]
        headers = ["S.No.", "Supplier"] + sorted_months + ["Total (Net)"]
        for idx, h in enumerate(headers, start=1):
            cell = ws_months.cell(row=row_idx, column=idx, value=h)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center")
            cell.border = table_border
        row_idx += 1
            
        for r_idx, row in enumerate(sorted(merged_data.values(), key=lambda x: x["total"], reverse=True)):
            fill = zebra_fill if r_idx % 2 == 0 else None
            ws_months.cell(row=row_idx, column=1, value=r_idx + 1).border = table_border
            ws_months.cell(row=row_idx, column=2, value=row["supp"]).border = table_border
            if fill:
                ws_months.cell(row=row_idx, column=1).fill = fill
                ws_months.cell(row=row_idx, column=2).fill = fill
            col_idx = 3
            for m_key in sorted_months:
                c = ws_months.cell(row=row_idx, column=col_idx, value=flt(row["months"].get(m_key, 0))/1000000)
                c.number_format, c.border = '[$₹-en-IN] #,##0.00 "M"', table_border
                if fill: c.fill = fill
                col_idx += 1
            c_n = ws_months.cell(row=row_idx, column=col_idx, value=flt(row["total"])/1000000)
            c_n.number_format = '[$₹-en-IN] #,##0.00 "M"'
            c_n.font = Font(bold=True)
            c_n.fill = PatternFill(start_color="ecf0f1", fill_type="solid")
            c_n.border = table_border
            col_idx += 1
            row_idx += 1
    
        ws_months.cell(row=row_idx, column=1, value="Grand Total").font = header_font
        ws_months.merge_cells(start_row=row_idx, start_column=1, end_row=row_idx, end_column=2)
        for c in range(1, 3):
            ws_months.cell(row=row_idx, column=c).fill = header_fill
            ws_months.cell(row=row_idx, column=c).border = table_border
            
        m_totals_net = {}
        g_total_net = sum(r["total"] for r in merged_data.values())
        
        for r in merged_data.values():
            for mk, mv in r["months"].items(): m_totals_net[mk] = m_totals_net.get(mk, 0) + mv
    
        col_idx = 3
        for m_key in sorted_months:
            c = ws_months.cell(row=row_idx, column=col_idx, value=flt(m_totals_net.get(m_key, 0))/1000000)
            c.number_format = '[$₹-en-IN] #,##0.00 "M"'
            c.font = header_font
            c.fill = header_fill
            c.border = table_border
            col_idx += 1
        c_gn = ws_months.cell(row=row_idx, column=col_idx, value=g_total_net / 1000000)
        c_gn.number_format = '[$₹-en-IN] #,##0.00 "M"'
        c_gn.font = header_font
        c_gn.fill = header_fill
        c_gn.border = table_border
        
        ws_months.column_dimensions["A"].width = 8
        ws_months.column_dimensions["B"].width = 35
        for i in range(3, 3 + len(sorted_months) + 1):
            ws_months.column_dimensions[get_column_letter(i)].width = 22

    if export_type in ["all", "detail"]:
        ws = wb.get_sheet_by_name("Supplier Orders List") if "Supplier Orders List" in wb.sheetnames else wb.create_sheet("Supplier Orders List")
        row_idx = 1
        ws.cell(row=row_idx, column=1, value="Supplier Order List").font = section_font
        row_idx += 2
    
        ui_columns = [
            {"label": "S.No.", "fieldname": "sr_no_idx", "width": 8},
            {"label": "PO No", "fieldname": "name", "width": 18},
            {"label": "Supplier", "fieldname": "supplier", "width": 35},
            {"label": "PO Date", "fieldname": "transaction_date", "width": 14},
            {"label": "Expected Del.", "fieldname": "schedule_date", "width": 14},
            {"label": "Agreed Time", "fieldname": "supplier_agreed_time", "width": 14},
            {"label": "Delivery as per PO", "fieldname": "delivery_time_as_per_po", "width": 14},
            {"label": "Actual Delivery", "fieldname": "actual_delivery_time", "width": 14},
            {"label": "Status", "fieldname": "status", "width": 18},
            {"label": "Net Total (M)", "fieldname": "net_total", "width": 18},
        ]
    
        for idx, col in enumerate(ui_columns, start=1):
            cell = ws.cell(row=row_idx, column=idx, value=col["label"])
            cell.font = header_font; cell.fill = header_fill; cell.alignment = Alignment(horizontal="center"); cell.border = table_border
            ws.column_dimensions[get_column_letter(idx)].width = col["width"]
        row_idx += 1
    
        total_amt = 0
    
        for r_idx, row in enumerate(data):
            row_fill = zebra_fill if r_idx % 2 != 0 else None
            for idx, col in enumerate(ui_columns, start=1):
                fname = col["fieldname"]
                val = row.get(fname)
                if fname == "sr_no_idx": val = r_idx + 1
                cell = ws.cell(row=row_idx, column=idx)
                cell.border = table_border
                if row_fill: cell.fill = row_fill
                
                if fname == "net_total":
                    num_val = flt(val or 0) / 1000000
                    total_amt += flt(val or 0)
                    cell.number_format = '[$₹-en-IN] #,##0.00 "M"'
                    cell.value = num_val; cell.alignment = Alignment(horizontal="right")
                elif fname in ["transaction_date", "schedule_date", "supplier_agreed_time", "delivery_time_as_per_po", "actual_delivery_time"]:
                    cell.value = str(val) if val else "-"; cell.alignment = Alignment(horizontal="center")
                else:
                    cell.value = str(val) if val is not None else ""; cell.alignment = Alignment(horizontal="left")
            row_idx += 1
    
        # Add Total Row
        c_tot_label = ws.cell(row=row_idx, column=1, value="GRAND TOTAL")
        c_tot_label.font = Font(bold=True)
        ws.merge_cells(start_row=row_idx, start_column=1, end_row=row_idx, end_column=9)
        for c in range(1, 10):
            ws.cell(row=row_idx, column=c).fill = header_fill
            ws.cell(row=row_idx, column=c).font = header_font
            ws.cell(row=row_idx, column=c).border = table_border
            if c == 1:
                ws.cell(row=row_idx, column=c).alignment = Alignment(horizontal="right", vertical="center")
    
        c_tot_amt = ws.cell(row=row_idx, column=10, value=total_amt / 1000000)
        c_tot_amt.font = header_font; c_tot_amt.fill = header_fill; c_tot_amt.border = table_border; c_tot_amt.alignment = Alignment(horizontal="right")
        c_tot_amt.number_format = '[$₹-en-IN] #,##0.00 "M"'

    # Remove dummy sheets if created
    for dummy_name in ["Dummy1", "Dummy2"]:
        if dummy_name in wb.sheetnames:
            wb.remove(wb[dummy_name])

    output = BytesIO()
    wb.save(output)
    output.seek(0)
    
    filename = f"Supplier_Order_Dashboard_{nowdate()}.xlsx"
    if export_type == "summary":
        filename = f"Supplier_Consolidated_Booking_{nowdate()}.xlsx"
    elif export_type == "detail":
        filename = f"Supplier_Orders_List_{nowdate()}.xlsx"

    return {
        "filename": filename,
        "filecontent": base64.b64encode(output.read()).decode()
    }
