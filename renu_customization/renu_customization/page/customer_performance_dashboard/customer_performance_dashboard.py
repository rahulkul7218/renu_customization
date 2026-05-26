import frappe
from frappe import _
from frappe.utils import flt, getdate, nowdate, add_days, formatdate
import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
from io import BytesIO
import base64

@frappe.whitelist()
def export_to_pdf(html=None, orientation="Landscape"):
    if not html:
        frappe.throw(_("PDF content is empty"))

    options = {
        "page-size": "A4",
        "orientation": orientation,
        "margin-top": "10mm",
        "margin-right": "10mm",
        "margin-bottom": "10mm",
        "margin-left": "10mm",
        "encoding": "UTF-8",
        "no-outline": None,
    }

    pdf_content = frappe.utils.pdf.get_pdf(html, options)

    frappe.local.response.filename = f"Customer_Performance_{nowdate()}.pdf"
    frappe.local.response.filecontent = pdf_content
    frappe.local.response.type = "download"

def _parse_group_concat_dates(dates_str):
    if not dates_str:
        return []
    return [getdate(d.strip()) for d in str(dates_str).split(",") if d and str(d).strip()]


def get_actual_delivery_dates(row):
    dates = row.get("actual_delivery_dates") or []
    if isinstance(dates, str):
        dates = _parse_group_concat_dates(dates)
    elif dates:
        dates = [getdate(d) for d in dates if d]
    if not dates and row.get("actual_delivery_time"):
        dates = [getdate(row["actual_delivery_time"])]
    return sorted(dates)


def format_actual_delivery_display(row):
    dates = get_actual_delivery_dates(row)
    if not dates:
        return None
    return ", ".join(formatdate(d) for d in dates)


def _fetch_actual_delivery_map(so_item_field):
    """Return {so_item_name: [date, ...]} from all submitted delivery notes."""
    actual_delivery_map = {}
    dn_data = frappe.db.sql(
        f"""
        SELECT
            dni.{so_item_field} AS so_item_name,
            GROUP_CONCAT(DISTINCT dn.posting_date ORDER BY dn.posting_date SEPARATOR ',') AS actual_delivery_dates
        FROM
            `tabDelivery Note` dn
        INNER JOIN
            `tabDelivery Note Item` dni ON dni.parent = dn.name
        WHERE
            dn.docstatus = 1
            AND dni.{so_item_field} IS NOT NULL
            AND dni.{so_item_field} != ''
        GROUP BY
            dni.{so_item_field}
        """,
        as_dict=True,
    )
    for r in dn_data:
        actual_delivery_map[r.so_item_name] = _parse_group_concat_dates(r.actual_delivery_dates)
    return actual_delivery_map


EXCLUDED_SO_STATUSES = ("Cancelled", "Draft")


def get_line_net_amount(item):
    short_close = flt(item.get("total_short_close_qty"))
    effective_qty = flt(item.get("qty")) - short_close
    base_rate = flt(item.get("base_rate"))
    if base_rate:
        return effective_qty * base_rate
    return flt(item.get("base_amount")) or flt(item.get("amount"))


def is_open_order_line(pending_qty, net_total):
    """Fully short-closed lines have no pending qty and no booked value."""
    return flt(pending_qty) > 0 or flt(net_total) > 0


def _get_freight_item_codes(item_codes):
    if not item_codes:
        return set()
    return set(
        frappe.db.sql_list(
            """
            SELECT name FROM `tabItem`
            WHERE name IN %s
            AND (IFNULL(custom_is_freight_item, 0) = 1 OR item_name = 'Freight')
            """,
            (tuple(item_codes),),
        )
    )


def prepare_filters(filters):
    if not filters:
        filters = {}
    elif isinstance(filters, str):
        filters = frappe.parse_json(filters)
    
    if filters.get("fiscal_year") and (not filters.get("from_date") or not filters.get("to_date")):
        fy = frappe.get_doc("Fiscal Year", filters.get("fiscal_year"))
        if fy:
            filters["from_date"] = filters.get("from_date") or fy.year_start_date
            filters["to_date"] = filters.get("to_date") or fy.year_end_date

    return frappe._dict(filters)

@frappe.whitelist()
def get_dashboard_data(filters=None):
    if not frappe.has_permission("Sales Order", "read"):
        frappe.throw(_("Not permitted to read Sales Order"), frappe.PermissionError)

    filters = prepare_filters(filters)

    so_filters = {"docstatus": 1}

    if filters.get("company"):
        so_filters["company"] = filters.get("company")

    if filters.get("from_date"):
        so_filters["transaction_date"] = [">=", filters.get("from_date")]

    if filters.get("to_date"):
        if "transaction_date" in so_filters:
            so_filters["transaction_date"] = ["between", [filters.get("from_date"), filters.get("to_date")]]
        else:
            so_filters["transaction_date"] = ["<=", filters.get("to_date")]

    if filters.get("sales_order"):
        so_filters["name"] = filters.get("sales_order")

    if filters.get("customer"):
        so_filters["customer"] = filters.get("customer")
    elif filters.get("customer_group"):
        try:
            lft, rgt = frappe.db.get_value("Customer Group", filters.customer_group, ["lft", "rgt"])
            customers = frappe.db.sql_list(
                "SELECT name FROM `tabCustomer` WHERE customer_group IN "
                "(SELECT name FROM `tabCustomer Group` WHERE lft >= %s AND rgt <= %s)",
                (lft, rgt),
            )
        except Exception:
            customers = frappe.get_all(
                "Customer", filters={"customer_group": filters.customer_group}, pluck="name"
            )
        so_filters["customer"] = ["in", customers] if customers else ["in", [""]]

    if filters.get("status"):
        if filters.get("status") in EXCLUDED_SO_STATUSES:
            return {"summary": [], "results": []}
        so_filters["status"] = filters.get("status")
    else:
        so_filters["status"] = ["not in", list(EXCLUDED_SO_STATUSES)]

    sales_orders = frappe.get_all(
        "Sales Order",
        filters=so_filters,
        fields=["name", "customer", "customer_name", "transaction_date", "delivery_date", "status"],
        order_by="transaction_date desc",
        limit_page_length=0,
    )

    if not sales_orders:
        return {"summary": [], "results": []}

    so_names = [so.name for so in sales_orders]
    so_map = {so.name: so for so in sales_orders}

    soi_filters = {"parent": ["in", so_names]}
    if filters.get("expected_delivery_date"):
        soi_filters["delivery_date"] = filters.get("expected_delivery_date")

    so_items = frappe.get_all(
        "Sales Order Item",
        filters=soi_filters,
        fields=[
            "name", "parent", "item_code", "item_name", "qty", "rate", "amount",
            "base_rate", "base_amount", "delivered_qty", "billed_amt",
            "delivery_date", "total_short_close_qty",
        ],
        order_by="parent, idx",
        limit_page_length=0,
    )

    if not so_items:
        return {"summary": [], "results": []}

    freight_item_codes = _get_freight_item_codes(list({i.item_code for i in so_items if i.item_code}))

    actual_delivery_map = {}
    try:
        actual_delivery_map = _fetch_actual_delivery_map("so_detail")
    except Exception:
        try:
            actual_delivery_map = _fetch_actual_delivery_map("sales_order_item")
        except Exception:
            pass

    report_data = []
    for item in so_items:
        so = so_map.get(item.parent)
        if not so or so.status in EXCLUDED_SO_STATUSES:
            continue
        if item.item_code in freight_item_codes:
            continue

        qty = flt(item.qty)
        delivered_qty = flt(item.delivered_qty)
        short_close_qty = flt(item.total_short_close_qty)
        rate = flt(item.rate)
        billed_amt = flt(item.billed_amt)
        billed_qty = billed_amt / rate if rate > 0 else 0.0
        pending_qty = max(0, qty - short_close_qty - delivered_qty)
        net_total = get_line_net_amount(item)
        if not is_open_order_line(pending_qty, net_total):
            continue

        per_delivered = (delivered_qty / qty) * 100 if qty > 0 else 0.0
        per_billed = (billed_qty / qty) * 100 if qty > 0 else 0.0

        schedule_date = item.delivery_date or so.delivery_date
        delivery_dates = actual_delivery_map.get(item.name, [])

        report_data.append({
            "name": so.name,
            "customer": so.customer_name or so.customer,
            "transaction_date": so.transaction_date,
            "status": so.status,
            "so_item_name": item.name,
            "item_code": item.item_code,
            "item_name": item.item_name,
            "qty": qty,
            "rate": rate,
            "net_total": net_total,
            "short_close_qty": short_close_qty,
            "delivered_qty": delivered_qty,
            "pending_qty": pending_qty,
            "billed_qty": billed_qty,
            "per_delivered": per_delivered,
            "per_billed": per_billed,
            "schedule_date": schedule_date,
            "actual_delivery_dates": [str(d) for d in delivery_dates],
            "actual_delivery_time": str(delivery_dates[-1]) if delivery_dates else None,
        })

    if not report_data:
        return {"summary": [], "results": []}

    results = []
    total_amount = 0
    total_overdue = 0
    total_due_next_15_days = 0
    unique_sos = set()

    today = getdate(nowdate())
    next_15_days = add_days(today, 15)

    for row in report_data:
        is_overdue = False
        due_next_15_days_flag = False
        row["due_days"] = "-"

        if row.get("schedule_date") and row.get("status") not in ["Completed", "Closed", "Cancelled"]:
            delivery_date = getdate(row.get("schedule_date"))
            pending_delivery = flt(row.get("pending_qty")) > 0

            if delivery_date < today and pending_delivery:
                is_overdue = True
                row["due_days"] = (today - delivery_date).days
            elif today <= delivery_date <= next_15_days and pending_delivery:
                due_next_15_days_flag = True
                row["due_days"] = (delivery_date - today).days
            elif delivery_date > today and pending_delivery:
                row["due_days"] = (delivery_date - today).days

        row["is_overdue"] = is_overdue
        row["is_due_next_15_days"] = due_next_15_days_flag

        if filters.get("is_overdue") and not is_overdue:
            continue
        if filters.get("due_next_15_days") and not due_next_15_days_flag:
            continue

        unique_sos.add(row.get("name"))
        total_amount += flt(row.get("net_total"))
        if is_overdue:
            total_overdue += flt(row.get("net_total"))
        if due_next_15_days_flag:
            total_due_next_15_days += flt(row.get("net_total"))

        results.append(row)

    results.sort(
        key=lambda x: getdate(x.get("schedule_date")) if x.get("schedule_date") else today,
        reverse=True,
    )

    processed_results = results

    summary = [
        {"label": _("Total Orders"), "value": len(unique_sos), "indicator": "blue", "fieldtype": "Int"},
        {"label": _("Total Net Amount"), "value": total_amount, "indicator": "green", "fieldtype": "Currency"},
        {"label": _("Overdue Amount"), "value": total_overdue, "indicator": "red", "fieldtype": "Currency"},
        {"label": _("Due Next 15 days"), "value": total_due_next_15_days, "indicator": "orange", "fieldtype": "Currency"}
    ]

    customer_totals = {}
    status_counts = {}
    
    for row in processed_results:
        cust = row.get("customer") or "Unknown"
        stat = row.get("status") or "Unknown"
        customer_totals[cust] = customer_totals.get(cust, 0) + flt(row.get("net_total"))
        status_counts[stat] = status_counts.get(stat, 0) + flt(row.get("net_total"))

    top_10_customers = sorted(customer_totals.items(), key=lambda x: x[1], reverse=True)[:10]
    
    charts = {
        "top_10_customers": {
            "title": _("Top 10 Customers"),
            "data": {
                "labels": [x[0] for x in top_10_customers],
                "datasets": [{"name": "Amount", "values": [x[1] for x in top_10_customers]}]
            },
            "type": "donut",
            "colors": ["#3b82f6", "#8b5cf6", "#10b981", "#f59e0b", "#06b6d4", "#ef4444", "#6366f1", "#ec4899", "#84cc16", "#f97316"],
            "is_currency": True
        },
        "order_status": {
            "title": _("Order Status Wise Amount "),
            "data": {
                "labels": list(status_counts.keys()),
                "datasets": [{"name": "Amount", "values": list(status_counts.values())}]
            },
            "type": "donut",
            "colors": ["#10b981", "#f59e0b", "#3b82f6", "#ef4444", "#8b5cf6", "#6366f1", "#ec4899", "#84cc16"],
            "is_currency": True
        }
    }

    # Monthly Calculations
    months_set = set()
    monthly_lifecycle = {
        "Booked": {},
        "Delivered": {},
        "Pending": {},
        "Overdue": {}
    }
    month_wise_customer = {}

    for row in processed_results:
        try:
            d = getdate(row.get("transaction_date"))
            m_key = d.strftime("%b %Y")
            m_sort = d.strftime("%Y%m")
            months_set.add((m_sort, m_key))
        except:
            m_key = "Unknown"
        
        amt = flt(row.get("net_total"))
        cust = row.get("customer") or "Unknown"

        # Lifecycle
        monthly_lifecycle["Booked"][m_key] = monthly_lifecycle["Booked"].get(m_key, 0) + amt
        if row.get("status") == "Completed":
            monthly_lifecycle["Delivered"][m_key] = monthly_lifecycle["Delivered"].get(m_key, 0) + amt
        
        # We'll calculate Pending and Overdue after processing all rows if needed, 
        # or just use the current status for simplified dashboard view
        if (
            flt(row.get("pending_qty")) > 0
            and row.get("status") not in ["Completed", "Closed", "Cancelled"]
        ):
            monthly_lifecycle["Pending"][m_key] = monthly_lifecycle["Pending"].get(m_key, 0) + amt
            if row.get("is_overdue"):
                monthly_lifecycle["Overdue"][m_key] = monthly_lifecycle["Overdue"].get(m_key, 0) + amt

        # Month-wise Customer Breakdown
        if cust not in month_wise_customer:
            month_wise_customer[cust] = {"customer": cust, "months": {}, "total": 0}
        month_wise_customer[cust]["months"][m_key] = month_wise_customer[cust]["months"].get(m_key, 0) + amt
        month_wise_customer[cust]["total"] += amt

    due_next_15_days = [row for row in processed_results if row.get("is_due_next_15_days")]

    sorted_months = [{"key": x[1], "sort": x[0]} for x in sorted(list(months_set), key=lambda x: x[0])]

    return {
        "summary": summary,
        "results": processed_results,
        "charts": charts,
        "months": sorted_months,
        "monthly_lifecycle": monthly_lifecycle,
        "month_wise_customer": sorted(month_wise_customer.values(), key=lambda x: x["total"], reverse=True),
        "due_next_15_days": sorted(due_next_15_days, key=lambda x: x.get("schedule_date"))
    }

def _customer_excel_styles():
    thin_side = Side(style="thin")
    return {
        "header_font": Font(bold=True, color="FFFFFF"),
        "header_fill": PatternFill(start_color="2c3e50", fill_type="solid"),
        "title_font": Font(bold=True, size=14),
        "section_font": Font(bold=True, size=12),
        "table_border": Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side),
        "zebra_fill": PatternFill(start_color="f8f9fa", fill_type="solid"),
        "total_font": Font(bold=True, size=11),
        "total_fill": PatternFill(start_color="e2e8f0", fill_type="solid"),
    }


def _write_customer_headers(ws, headers, styles, row_idx=1):
    for idx, header in enumerate(headers, start=1):
        cell = ws.cell(row=row_idx, column=idx, value=header)
        cell.font = styles["header_font"]
        cell.fill = styles["header_fill"]
        cell.alignment = Alignment(horizontal="center")
        cell.border = styles["table_border"]


def _write_customer_row(ws, row_idx, values, styles, amount_col=None):
    row_fill = styles["zebra_fill"] if row_idx % 2 == 0 else None
    for c_idx, val in enumerate(values, start=1):
        cell = ws.cell(row=row_idx, column=c_idx, value=val)
        cell.border = styles["table_border"]
        if row_fill:
            cell.fill = row_fill
        if amount_col and c_idx == amount_col:
            cell.number_format = "#,##0.00"
            cell.alignment = Alignment(horizontal="right")


def _write_customer_total_row(ws, row_idx, label, amount_col, amount_value, styles):
    label_end_col = max(1, amount_col - 1)
    for col in range(1, amount_col + 1):
        cell = ws.cell(row=row_idx, column=col)
        cell.fill = styles["total_fill"]
        cell.border = styles["table_border"]
    if label_end_col > 1:
        ws.merge_cells(start_row=row_idx, start_column=1, end_row=row_idx, end_column=label_end_col)
    label_cell = ws.cell(row=row_idx, column=1, value=label)
    label_cell.font = styles["total_font"]
    label_cell.fill = styles["total_fill"]
    label_cell.border = styles["table_border"]
    label_cell.alignment = Alignment(horizontal="right", vertical="center")
    amount_cell = ws.cell(row=row_idx, column=amount_col, value=amount_value)
    amount_cell.font = styles["total_font"]
    amount_cell.fill = styles["total_fill"]
    amount_cell.border = styles["table_border"]
    amount_cell.number_format = "#,##0.00"
    amount_cell.alignment = Alignment(horizontal="right")


def _write_customer_chart_tables(ws, dashboard_data, styles, start_row):
    charts = dashboard_data.get("charts") or {}
    row_idx = start_row
    chart_headers = ["Label", "Amount (M)", "% Share"]

    for chart_obj in charts.values():
        title = chart_obj.get("title") or "Chart"
        labels = chart_obj.get("data", {}).get("labels") or []
        values = (chart_obj.get("data", {}).get("datasets") or [{}])[0].get("values") or []
        is_currency = chart_obj.get("is_currency")
        chart_total = sum(flt(v) for v in values)

        ws.cell(row=row_idx, column=1, value=title).font = styles["section_font"]
        row_idx += 1
        _write_customer_headers(ws, chart_headers, styles, row_idx=row_idx)
        row_idx += 1

        for label, val in zip(labels, values):
            amount = flt(val) / 1000000 if is_currency else flt(val)
            pct = (flt(val) / chart_total * 100) if chart_total else 0
            _write_customer_row(ws, row_idx, [label, amount, round(pct, 1)], styles, amount_col=2)
            row_idx += 1

        _write_customer_row(
            ws, row_idx, ["TOTAL", chart_total / 1000000 if is_currency else chart_total, 100],
            styles, amount_col=2,
        )
        for col in range(1, 4):
            cell = ws.cell(row=row_idx, column=col)
            cell.font = styles["total_font"]
            cell.fill = styles["total_fill"]
        row_idx += 2

    return row_idx


def _write_customer_section_title(ws, row_idx, title, styles):
    ws.cell(row=row_idx, column=1, value=title).font = styles["section_font"]
    return row_idx + 1


def _autofit_customer_sheet(ws):
    for col in ws.columns:
        max_length = 0
        column = col[0].column_letter
        for cell in col:
            try:
                if cell.value is not None:
                    max_length = max(max_length, len(str(cell.value)))
            except Exception:
                pass
        ws.column_dimensions[column].width = min(max_length + 3, 55)


def _write_customer_overview_sheet(ws, dashboard_data, styles, results=None, due_rows=None):
    summary = dashboard_data.get("summary") or []
    ws.cell(row=1, column=1, value="Customer Performance Dashboard Overview").font = styles["title_font"]
    ws.cell(row=1, column=4, value="Generated On: " + str(nowdate()))
    row_idx = 3
    row_idx = _write_customer_section_title(ws, row_idx, "Operational Summary (Million INR)", styles)
    colors = {"blue": "3498db", "green": "2ecc71", "red": "e74c3c", "orange": "e67e22"}
    for i, s in enumerate(summary):
        r = row_idx + (i // 4) * 3
        c = 1 + (i % 4) * 2
        bg_color = colors.get((s.get("indicator") or "blue").lower(), "3498db")
        cell_l = ws.cell(row=r, column=c, value=s.get("label"))
        cell_l.font = Font(bold=True, color="FFFFFF")
        cell_l.fill = PatternFill(start_color=bg_color, fill_type="solid")
        cell_l.alignment = Alignment(horizontal="center")
        val = flt(s.get("value"))
        if s.get("fieldtype") == "Currency":
            val = val / 1000000
        cell_v = ws.cell(row=r + 1, column=c, value=val)
        cell_v.font = Font(bold=True, size=12)
        if s.get("fieldtype") == "Currency":
            cell_v.number_format = '"₹ "#,##0.00" M"'
        cell_v.alignment = Alignment(horizontal="center")
        ws.merge_cells(start_row=r, start_column=c, end_row=r, end_column=c + 1)
        ws.merge_cells(start_row=r + 1, start_column=c, end_row=r + 1, end_column=c + 1)

    row_idx = row_idx + 4
    row_idx = _write_customer_section_title(ws, row_idx, "Visual Analytics", styles)
    row_idx = _write_customer_chart_tables(ws, dashboard_data, styles, row_idx)

    if results is not None and due_rows is not None:
        row_idx += 1
        row_idx = _write_customer_section_title(ws, row_idx, "Month-Wise Booking Breakdown ", styles)
        row_idx = _write_customer_month_table(ws, dashboard_data, styles, row_idx)

        row_idx += 1
        row_idx = _write_customer_section_title(ws, row_idx, "Orders Due in Next 15 Days ", styles)
        row_idx = _write_customer_due_table(ws, due_rows, styles, row_idx)

        row_idx += 1
        row_idx = _write_customer_section_title(ws, row_idx, "Detailed Customer Orders List ", styles)
        _write_customer_detail_table(ws, results, styles, row_idx)

    _autofit_customer_sheet(ws)


def _write_customer_month_table(ws, dashboard_data, styles, start_row=1):
    months = dashboard_data.get("months") or []
    headers = ["S.No.", "Customer"] + [m["key"] for m in months] + ["Total (M)"]
    _write_customer_headers(ws, headers, styles, row_idx=start_row)
    month_rows = dashboard_data.get("month_wise_customer") or []
    month_totals = {m["key"]: 0 for m in months}
    grand_total = 0
    data_start = start_row + 1
    for i, row in enumerate(month_rows):
        r_idx = data_start + i
        vals = [i + 1, row.get("customer")]
        for m in months:
            amt = flt(row.get("months", {}).get(m["key"], 0))
            month_totals[m["key"]] += amt
            vals.append(amt / 1000000)
        row_total = flt(row.get("total"))
        grand_total += row_total
        vals.append(row_total / 1000000)
        _write_customer_row(ws, r_idx, vals, styles, amount_col=len(vals))
    total_row_idx = data_start + len(month_rows)
    ws.merge_cells(start_row=total_row_idx, start_column=1, end_row=total_row_idx, end_column=2)
    total_label = ws.cell(row=total_row_idx, column=1, value="GRAND TOTAL")
    total_label.font = styles["total_font"]
    total_label.fill = styles["total_fill"]
    total_label.border = styles["table_border"]
    total_label.alignment = Alignment(horizontal="right")
    ws.cell(row=total_row_idx, column=2).fill = styles["total_fill"]
    ws.cell(row=total_row_idx, column=2).border = styles["table_border"]
    for col_idx, m in enumerate(months, start=3):
        cell = ws.cell(row=total_row_idx, column=col_idx, value=month_totals[m["key"]] / 1000000)
        cell.font = styles["total_font"]
        cell.fill = styles["total_fill"]
        cell.border = styles["table_border"]
        cell.number_format = "#,##0.00"
        cell.alignment = Alignment(horizontal="right")
    grand_cell = ws.cell(row=total_row_idx, column=len(headers), value=grand_total / 1000000)
    grand_cell.font = styles["total_font"]
    grand_cell.fill = styles["total_fill"]
    grand_cell.border = styles["table_border"]
    grand_cell.number_format = "#,##0.00"
    grand_cell.alignment = Alignment(horizontal="right")
    if start_row == 1:
        _autofit_customer_sheet(ws)
    return total_row_idx + 1


def _write_customer_month_sheet(ws, dashboard_data, styles):
    _write_customer_month_table(ws, dashboard_data, styles, start_row=1)


def _customer_order_row_values(row, serial_no, include_days_left=False):
    pending_qty = flt(row.get("pending_qty"))
    if not pending_qty and row.get("qty") is not None:
        pending_qty = max(0, flt(row.get("qty")) - flt(row.get("delivered_qty")))

    vals = [
        serial_no,
        row.get("name"),
        row.get("customer"),
        row.get("item_code"),
        row.get("item_name"),
        row.get("transaction_date"),
        row.get("schedule_date"),
        format_actual_delivery_display(row),
    ]
    if include_days_left:
        due_days = row.get("due_days")
        vals.append(f"{due_days} Days" if due_days != "-" else due_days)
    vals.extend([
        row.get("status"),
        int(row.get("per_delivered") or 0),
        int(row.get("per_billed") or 0),
        flt(row.get("qty")),
        flt(row.get("delivered_qty")),
        pending_qty,
        flt(row.get("net_total")) / 1000000,
    ])
    return vals


CUSTOMER_DUE_HEADERS = [
    "S.No.", "SO No", "Customer", "Item Code", "Item Name", "Order Date",
    "Expected Del.", "Actual Del.", "Days Left", "Status", "% Del.", "% Bill.",
    "Order Qty", "Delivered Qty", "Pending Qty", "Net Total (M)",
]

CUSTOMER_DETAIL_HEADERS = [
    "S.No.", "SO No", "Customer", "Item Code", "Item Name", "Order Date",
    "Expected Del.", "Actual Del.", "Status", "% Del.", "% Bill.",
    "Order Qty", "Delivered Qty", "Pending Qty", "Net Total (M)",
]


def _customer_due_row_values(row, serial_no):
    return _customer_order_row_values(row, serial_no, include_days_left=True)


def _customer_detail_row_values(row, serial_no):
    return _customer_order_row_values(row, serial_no, include_days_left=False)


def _write_customer_due_table(ws, rows, styles, start_row=1):
    _write_customer_headers(ws, CUSTOMER_DUE_HEADERS, styles, row_idx=start_row)
    data_rows = rows or []
    data_start = start_row + 1
    for i, row in enumerate(data_rows):
        r_idx = data_start + i
        _write_customer_row(
            ws, r_idx, _customer_due_row_values(row, i + 1), styles, amount_col=len(CUSTOMER_DUE_HEADERS),
        )
    total_amount = sum(flt(r.get("net_total")) for r in data_rows) / 1000000
    _write_customer_total_row(
        ws, data_start + len(data_rows), "TOTAL DUE VALUE", len(CUSTOMER_DUE_HEADERS), total_amount, styles,
    )
    if start_row == 1:
        _autofit_customer_sheet(ws)
    return data_start + len(data_rows) + 1


def _write_customer_due_sheet(ws, rows, styles):
    _write_customer_due_table(ws, rows, styles, start_row=1)


def _write_customer_detail_table(ws, rows, styles, start_row=1):
    _write_customer_headers(ws, CUSTOMER_DETAIL_HEADERS, styles, row_idx=start_row)
    data_rows = rows or []
    data_start = start_row + 1
    for i, row in enumerate(data_rows):
        r_idx = data_start + i
        _write_customer_row(
            ws, r_idx, _customer_detail_row_values(row, i + 1), styles, amount_col=len(CUSTOMER_DETAIL_HEADERS),
        )
    total_amount = sum(flt(r.get("net_total")) for r in data_rows) / 1000000
    _write_customer_total_row(
        ws, data_start + len(data_rows), "TOTAL BOOKED VALUE", len(CUSTOMER_DETAIL_HEADERS), total_amount, styles,
    )
    if start_row == 1:
        _autofit_customer_sheet(ws)
    return data_start + len(data_rows) + 1


def _write_customer_detail_sheet(ws, rows, styles):
    _write_customer_detail_table(ws, rows, styles, start_row=1)


@frappe.whitelist()
def export_to_excel(filters=None, export_type="all"):
    filters = prepare_filters(filters)
    dashboard_data = get_dashboard_data(filters)
    results = dashboard_data.get("results") or []
    due_rows = dashboard_data.get("due_next_15_days") or []

    if not results and not dashboard_data.get("month_wise_customer"):
        frappe.throw(_("No data to export"))

    styles = _customer_excel_styles()
    wb = openpyxl.Workbook()

    if export_type == "all":
        ws_overview = wb.active
        ws_overview.title = "Overview"
        _write_customer_overview_sheet(ws_overview, dashboard_data, styles)

        ws_months = wb.create_sheet("Month-Wise Booking")
        _write_customer_month_sheet(ws_months, dashboard_data, styles)

        ws_due = wb.create_sheet("Due in 15 Days")
        _write_customer_due_sheet(ws_due, due_rows, styles)

        ws_list = wb.create_sheet("Detailed Orders")
        _write_customer_detail_sheet(ws_list, results, styles)
    elif export_type == "summary":
        ws_months = wb.active
        ws_months.title = "Month-Wise Booking"
        _write_customer_month_sheet(ws_months, dashboard_data, styles)
    elif export_type == "due":
        ws_due = wb.active
        ws_due.title = "Due in 15 Days"
        _write_customer_due_sheet(ws_due, due_rows, styles)
    elif export_type == "detail":
        ws_list = wb.active
        ws_list.title = "Detailed Orders"
        _write_customer_detail_sheet(ws_list, results, styles)
    else:
        frappe.throw(_("Invalid export type"))

    output = BytesIO()
    wb.save(output)
    output.seek(0)

    filenames = {
        "all": f"Customer_Performance_{nowdate()}.xlsx",
        "summary": f"Customer_Performance_Month_Wise_{nowdate()}.xlsx",
        "due": f"Customer_Performance_Due_15_Days_{nowdate()}.xlsx",
        "detail": f"Customer_Performance_Detailed_{nowdate()}.xlsx",
    }

    return {
        "filename": filenames.get(export_type, filenames["all"]),
        "filecontent": base64.b64encode(output.getvalue()).decode(),
    }