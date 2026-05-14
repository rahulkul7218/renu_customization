import frappe
from frappe import _
from frappe.utils import flt, getdate
import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from io import BytesIO
import base64
import json

@frappe.whitelist()
def export_to_pdf(html):
	frappe.response.filename = "sales_revenue_dashboard.pdf"
	frappe.response.type = "download"
	frappe.response.filecontent = frappe.utils.pdf.get_pdf(html, {"orientation": "Landscape"})


def prepare_filters(filters):
    if not filters:
        filters = {}
    elif isinstance(filters, str):
        filters = frappe.parse_json(filters)
    
    filters = frappe._dict(filters)

    # Handle DateRange from JS (Legacy/Compatibility)
    if filters.get("date_range"):
        dr = filters.get("date_range")
        if isinstance(dr, list) and len(dr) == 2:
            filters["from_date"] = dr[0]
            filters["to_date"] = dr[1]

    # Handle Company Default
    if not filters.get("company"):
        filters["company"] = frappe.defaults.get_user_default("company") or \
                           frappe.db.get_single_value('Global Defaults', 'default_company')

    # Handle Fiscal Year Interaction
    if filters.get("fiscal_year"):
        fy = frappe.get_doc("Fiscal Year", filters.get("fiscal_year"))
        if fy:
            # 1. If NO dates provided, use FY defaults
            if not filters.get("from_date"):
                filters["from_date"] = fy.year_start_date
            if not filters.get("to_date"):
                filters["to_date"] = fy.year_end_date
            
            # 2. If dates ARE provided, constrain them to the FY bounds
            if filters.get("from_date") and getdate(filters.from_date) < getdate(fy.year_start_date):
                filters["from_date"] = fy.year_start_date
            
            if filters.get("to_date") and getdate(filters.to_date) > getdate(fy.year_end_date):
                filters["to_date"] = fy.year_end_date
    
    # Remove "All" values from filters so they don't affect backend queries
    keys_to_remove = [k for k, v in filters.items() if v == "All"]
    for k in keys_to_remove:
        del filters[k]

    return filters

@frappe.whitelist()
def get_dashboard_data(filters=None):
    filters = prepare_filters(filters)
    
    # 1. Base SQL Conditions
    conditions = "WHERE si.docstatus = 1 AND si.status != 'Cancelled'"
    query_params = {}

    # Standard Filters
    if filters.get("from_date"):
        conditions += " AND si.posting_date >= %(from_date)s"
        query_params["from_date"] = filters.from_date
    if filters.get("to_date"):
        conditions += " AND si.posting_date <= %(to_date)s"
        query_params["to_date"] = filters.to_date
    if filters.get("company"):
        conditions += " AND si.company = %(company)s"
        query_params["company"] = filters.company
    if filters.get("customer"):
        conditions += " AND si.customer = %(customer)s"
        query_params["customer"] = filters.customer
    if filters.get("customer_group"):
        conditions += " AND c.customer_group = %(customer_group)s"
        query_params["customer_group"] = filters.customer_group
    if filters.get("item_group"):
        conditions += " AND i.item_group = %(item_group)s"
        query_params["item_group"] = filters.item_group
    if filters.get("item_code"):
        conditions += " AND sii.item_code = %(item_code)s"
        query_params["item_code"] = filters.item_code
    if filters.get("sales_person"):
        conditions += " AND EXISTS (SELECT 1 FROM `tabSales Team` WHERE parent = si.name AND sales_person = %(sales_person)s)"
        query_params["sales_person"] = filters.sales_person
    if filters.get("business_region_name"):
        conditions += " AND c.business_region_name = %(business_region_name)s"
        query_params["business_region_name"] = filters.business_region_name
    if filters.get("dom_exp"):
        if filters.dom_exp == "Domestic":
            conditions += " AND (si.is_domestic = 1 OR IFNULL(si.is_export, 0) = 0)"
        elif filters.dom_exp == "Export":
            conditions += " AND si.is_export = 1"
    if filters.get("invoice_type"):
        conditions += " AND si.invoice_type = %(invoice_type)s"
        query_params["invoice_type"] = filters.invoice_type
    if filters.get("status"):
        if isinstance(filters.status, str):
            status_list = [s.strip() for s in filters.status.split(",") if s.strip()]
        else:
            status_list = filters.status
        conditions += " AND si.status IN %(status_list)s"
        query_params["status_list"] = tuple(status_list)

    # 2. Main Transactional Query
    # We join Sales Invoice Item with Sales Team to handle revenue splits correctly
    # We only include items that hit 'Income' root type accounts to match P&L
    sql = f"""
        SELECT 
            si.name as invoice_id,
            si.posting_date as invoice_date,
            si.customer,
            si.customer_name,
            si.status,
            si.invoice_type,
            si.is_domestic,
            si.is_export,
            si.base_net_total,
            si.base_grand_total,
            sii.item_code,
            sii.item_name,
            sii.base_net_amount as item_net_amount,
            sii.income_account,
            sii.qty,
            st.sales_person,
            st.allocated_percentage,
            c.customer_group,
            c.business_region_name,
            i.item_group,
            (SELECT SUM(base_tax_amount) 
             FROM `tabSales Taxes and Charges` 
             WHERE parent = si.name 
             AND account_head IN (SELECT name FROM `tabAccount` WHERE root_type = 'Income')) as tax_income
        FROM `tabSales Invoice` si
        JOIN `tabSales Invoice Item` sii ON sii.parent = si.name
        LEFT JOIN `tabSales Team` st ON st.parent = si.name
        LEFT JOIN `tabCustomer` c ON c.name = si.customer
        LEFT JOIN `tabItem` i ON i.name = sii.item_code
        JOIN `tabAccount` acc ON acc.name = sii.income_account AND acc.root_type = 'Income'
        {conditions} AND IFNULL(i.custom_is_freight_item, 0) = 0
        ORDER BY si.posting_date DESC, si.name DESC
    """
    
    raw_data = frappe.db.sql(sql, query_params, as_dict=1)
    
    if not raw_data:
        return { "summary": [], "charts": {}, "results": [], "columns": [] }

    # 3. Process Data & Calculate Metrics (Million INR)
    data = []
    total_net_rev = 0
    total_grand_rev = 0
    dom_rev = 0
    exp_rev = 0
    cp_rev = 0
    
    sp_revenue = {}
    cust_revenue = {}
    prod_revenue = {}
    
    processed_items = set()

    for row in raw_data:
        # Avoid duplicate counting for multi-salesperson rows in grand totals
        # But for SP charts, we must split the amount
        inv_id = row.invoice_id
        
        alloc_p = flt(row.allocated_percentage) or 100
        
        # Revenue Attribution: Match "Sales Invoice Data" exactly (Base Net Amount)
        # We do NOT add tax_income here to avoid double-counting or mismatch with standard reports.
        total_item_revenue = flt(row.item_net_amount)
        
        amt_allocated = total_item_revenue * (alloc_p / 100)
        qty_allocated = flt(row.qty) * (alloc_p / 100)
        
        # Gross Attribution (Proportionate share of Grand Total)
        si_net_total = flt(row.base_net_total) or 1
        gross_factor = (flt(row.base_grand_total) / si_net_total) if si_net_total else 1.0
        gross_amt_allocated = amt_allocated * gross_factor
        
        # Store processed row for table
        row_copy = frappe._dict(row)
        row_copy.amt_allocated = amt_allocated / 1000000
        row_copy.gross_amount = gross_amt_allocated / 1000000
        row_copy.qty = qty_allocated # Split qty for correct totals
        row_copy.dom_exp = "Export" if row.is_export else "Domestic"
        
        # Classification for Channel Partner
        is_cp = False
        if row.customer_group:
            cg = row.customer_group.lower()
            if any(term in cg for term in ["system integrator", "distributor", "partner", "reseller"]):
                is_cp = True
        row_copy.is_channel_partner = is_cp
        
        data.append(row_copy)

        # Aggregate for KPIs (Only once per item-SP combination)
        total_net_rev += amt_allocated
        total_grand_rev += gross_amt_allocated
        
        if row.is_export:
            exp_rev += amt_allocated
        else:
            dom_rev += amt_allocated
            
        if is_cp:
            cp_rev += amt_allocated

        # Aggregate for Charts
        sp = row.sales_person or "No Sales Person"
        sp_revenue[sp] = sp_revenue.get(sp, 0) + amt_allocated
        
        cust = row.customer_name or row.customer or "Unknown"
        # Since we have multiple rows per customer (per item), we should only add this once per item-SP
        # The logic above already handles this by iterating over every item row
        cust_revenue[cust] = cust_revenue.get(cust, 0) + amt_allocated
        
        prod = row.item_name or row.item_code or "Unknown"
        prod_revenue[prod] = prod_revenue.get(prod, 0) + amt_allocated

    # 4. Final Output Generation
    report_summary = [
        {"label": _("Net Revenue"), "value": total_net_rev / 1000000, "indicator": "blue", "fieldtype": "Currency", "currency": "INR"},
        {"label": _("Gross Revenue"), "value": total_grand_rev / 1000000, "indicator": "cyan", "fieldtype": "Currency", "currency": "INR"},
        {"label": _("Domestic Net Revenue"), "value": dom_rev / 1000000, "indicator": "green", "fieldtype": "Currency", "currency": "INR"},
        {"label": _("Export Net Revenue"), "value": exp_rev / 1000000, "indicator": "orange", "fieldtype": "Currency", "currency": "INR"},
        {"label": _("Channel Partner"), "value": cp_rev / 1000000, "indicator": "purple", "fieldtype": "Currency", "currency": "INR"}
    ]

    def get_chart_def(title, data_dict, limit=10):
        sorted_items = sorted(data_dict.items(), key=lambda x: x[1], reverse=True)
        top_items = sorted_items[:limit]
        return {
            "title": title,
            "data": {
                "labels": [x[0] for x in top_items],
                "datasets": [{"name": title, "values": [flt(x[1] / 1000000, 4) for x in top_items]}]
            },
            "type": "donut",
            "height": 300,
            "colors": ['#3498db', '#e74c3c', '#2ecc71', '#f1c40f', '#9b59b6', '#34495e', '#ecf0f1', '#1abc9c', '#d35400', '#7f8c8d']
        }

    return {
        "summary": report_summary,
        "charts": {
            "top_10_salesperson": get_chart_def("Top 10 Salesperson", sp_revenue, limit=10),
            "top_10_customers": get_chart_def("Top 10 Customers", cust_revenue, limit=10),
            "top_10_products": get_chart_def("Top 10 Products", prod_revenue, limit=10)
        },
        "results": data,
        "columns": [] # Frontend handles column rendering
    }

from openpyxl.styles import Alignment, Font, PatternFill, Border, Side

@frappe.whitelist()
def export_to_excel(filters=None, export_type="all"):
    filters = prepare_filters(filters)
    
    # Use the dashboard's data fetching logic to ensure filters are applied
    dashboard_data = get_dashboard_data(filters)
    data = dashboard_data.get("results")
    summary = dashboard_data.get("summary")
    charts = dashboard_data.get("charts")
    
    if not data:
        return None

    wb = openpyxl.Workbook()
    
    # Conditionally create sheets based on export_type
    if export_type == "all":
        ws_overview = wb.active
        ws_overview.title = "Dashboard Overview"
        ws_months = wb.create_sheet("Month-Wise Revenue")
        ws_list = wb.create_sheet("Sales Invoices List")
    elif export_type == "summary":
        ws_months = wb.active
        ws_months.title = "Month-Wise Revenue"
        # We might need dummy sheets for the logic below if we don't wrap it carefully
        ws_overview = wb.create_sheet("Dummy1")
        ws_list = wb.create_sheet("Dummy2")
    elif export_type == "detail":
        ws_list = wb.active
        ws_list.title = "Sales Invoices List"
        ws_overview = wb.create_sheet("Dummy1")
        ws_months = wb.create_sheet("Dummy2")
    
    # Styling Helpers
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="2c3e50", fill_type="solid")
    title_font = Font(bold=True, size=14)
    section_font = Font(bold=True, size=12)
    footer_font = Font(bold=True, color="000000")
    thin_side = Side(style='thin')
    table_border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)
    zebra_fill = PatternFill(start_color="f8f9fa", fill_type="solid")

    row_idx = 1
    
    # -------------------------------------------------------------------------
    # Sheet 1: Dashboard Overview (KPIs & Top 10 Tables)
    # -------------------------------------------------------------------------
    if export_type == "all":
        ws_overview.cell(row=row_idx, column=1, value="Sales Revenue Dashboard Overview").font = title_font
        ws_overview.cell(row=row_idx, column=5, value="Generated On: " + frappe.utils.now_datetime().strftime("%Y-%m-%d %H:%M:%S"))
        row_idx += 2
     
        # 1. Summary Section (KPI Grid)
        ws_overview.cell(row=row_idx, column=1, value="1. Revenue Summary (Million INR)").font = section_font
        row_idx += 1
        
        summary_start_row = row_idx
        colors = {
            "blue": "3498db", "green": "2ecc71", "orange": "e67e22", "purple": "9b59b6", "red": "e74c3c"
        }
     
        for i, s in enumerate(summary):
            r = summary_start_row + (i // 4) * 3
            c = 1 + (i % 4) * 2
            
            # Label
            cell_l = ws_overview.cell(row=r, column=c, value=s.get('label'))
            cell_l.font = Font(bold=True, color="FFFFFF")
            indicator = s.get('indicator', 'blue').lower()
            bg_color = colors.get(indicator, "3498db")
            cell_l.fill = PatternFill(start_color=bg_color, fill_type="solid")
            cell_l.alignment = Alignment(horizontal="center")
            
            # Value (Converted to Million)
            val = flt(s.get('value'))
            cell_v = ws_overview.cell(row=r+1, column=c, value=val)
            cell_v.font = Font(bold=True, size=12)
            cell_v.number_format = '"₹ "#,##0.00" M"'
            cell_v.alignment = Alignment(horizontal="center")
            cell_v.border = Border(left=Side(style='medium', color=bg_color), 
                                   right=Side(style='medium', color=bg_color), 
                                   bottom=Side(style='medium', color=bg_color))
            
            ws_overview.merge_cells(start_row=r, start_column=c, end_row=r, end_column=c+1)
            ws_overview.merge_cells(start_row=r+1, start_column=c, end_row=r+1, end_column=c+1)
     
        row_idx = summary_start_row + 6
        row_idx += 1
     
        def write_chart_section(title, chart_data):
            nonlocal row_idx
            ws_overview.cell(row=row_idx, column=1, value=title).font = section_font
            row_idx += 1
            
            headers = ["Category", "Amount (M)", "Share %"]
            for idx, h in enumerate(headers, start=1):
                cell = ws_overview.cell(row=row_idx, column=idx, value=h)
                cell.font = header_font
                cell.fill = PatternFill(start_color="34495e", fill_type="solid")
                cell.border = table_border
                cell.alignment = Alignment(horizontal="center")
            row_idx += 1
                
            labels = chart_data.get("data", {}).get("labels", [])
            values = chart_data.get("data", {}).get("datasets", [{}])[0].get("values", [])
            total_val = sum(values) if values else 1
            
            for i in range(len(labels)):
                row_fill = zebra_fill if i % 2 == 0 else None
                c1 = ws_overview.cell(row=row_idx, column=1, value=labels[i])
                c1.border = table_border
                if row_fill: c1.fill = row_fill
                
                val_m = flt(values[i])
                c2 = ws_overview.cell(row=row_idx, column=2, value=val_m)
                c2.number_format = '"₹ "#,##0.00" M"'
                c2.border = table_border
                c2.alignment = Alignment(horizontal="right")
                if row_fill: c2.fill = row_fill
                
                share = (values[i] / total_val) if total_val else 0
                c3 = ws_overview.cell(row=row_idx, column=3, value=share)
                c3.number_format = "0.00%"
                c3.border = table_border
                c3.alignment = Alignment(horizontal="center")
                if row_fill: c3.fill = row_fill
                row_idx += 1
            row_idx += 2
            
        if charts.get("top_10_salesperson", {}).get("data", {}).get("labels"):
            write_chart_section("2. Top 10 Salesperson", charts["top_10_salesperson"])
        if charts.get("top_10_customers", {}).get("data", {}).get("labels"):
            write_chart_section("3. Top 10 Customers", charts["top_10_customers"])
        if charts.get("top_10_products", {}).get("data", {}).get("labels"):
            write_chart_section("4. Top 10 Products", charts["top_10_products"])

    if export_type in ["all", "summary"]:
        # -------------------------------------------------------------------------
        # Sheet 2: Month-Wise Revenue
        # -------------------------------------------------------------------------
        row_idx = 1
        ws_months.cell(row=row_idx, column=1, value="Month-Wise Consolidated Revenue (Million INR)").font = section_font
        row_idx += 1
        
        merged_data = {}
        months_set = set()
        for row in data:
            sp = row.get("sales_person") or "-"
            cust = row.get("customer_name") or row.get("customer") or "-"
            prod = row.get("item_name") or row.get("item_code") or "-"
            amt = row.get("amt_allocated") or 0
            gross_amt = row.get("gross_amount") or amt
    
            date_str = str(row.get("delivery_date") or row.get("invoice_date") or row.get("posting_date") or "")
            try:
                d = frappe.utils.getdate(date_str)
                m_key = d.strftime("%b %Y")
                m_sort = d.strftime("%Y%m")
            except:
                m_key = "Unknown"
                m_sort = "000000"
                
            months_set.add((m_sort, m_key))
            key = f"{cust}|{sp}|{prod}"
            if key not in merged_data:
                merged_data[key] = {"cust": cust, "sp": sp, "prod": prod, "months": {}, "total": 0, "total_gross": 0}
            merged_data[key]["months"][m_key] = merged_data[key]["months"].get(m_key, 0) + amt
            merged_data[key]["total"] += amt
            merged_data[key]["total_gross"] += gross_amt
            
        sorted_months = [x[1] for x in sorted(list(months_set), key=lambda x: x[0])]
        headers = ["S.No.", "Customer", "Sales Person", "Product"] + sorted_months + ["Total (Net)", "Grand Total (Gross)"]
        
        for idx, h in enumerate(headers, start=1):
            cell = ws_months.cell(row=row_idx, column=idx, value=h)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center")
            cell.border = table_border
        row_idx += 1
            
        for r_idx, row in enumerate(sorted(merged_data.values(), key=lambda x: x["total"], reverse=True)):
            row_fill = zebra_fill if r_idx % 2 != 0 else None
            
            c_no = ws_months.cell(row=row_idx, column=1, value=r_idx + 1)
            c1 = ws_months.cell(row=row_idx, column=2, value=row["cust"])
            c2 = ws_months.cell(row=row_idx, column=3, value=row["sp"])
            c3 = ws_months.cell(row=row_idx, column=4, value=row["prod"])
            for c in [c_no, c1, c2, c3]:
                c.border = table_border
                if row_fill: c.fill = row_fill
                
            col_idx = 5
            for m_key in sorted_months:
                v_m = flt(row["months"].get(m_key, 0))
                c = ws_months.cell(row=row_idx, column=col_idx, value=v_m)
                c.number_format = '"₹ "#,##0.00" M"'
                c.border = table_border
                c.alignment = Alignment(horizontal="right")
                if row_fill: c.fill = row_fill
                col_idx += 1
                
            tot_m = flt(row["total"])
            c_tot = ws_months.cell(row=row_idx, column=col_idx, value=tot_m)
            c_tot.number_format = '"₹ "#,##0.00" M"'
            c_tot.font = Font(bold=True)
            c_tot.fill = PatternFill(start_color="ecf0f1", fill_type="solid")
            c_tot.border = table_border
            c_tot.alignment = Alignment(horizontal="right")
            col_idx += 1
            
            tot_g = flt(row["total_gross"])
            c_g = ws_months.cell(row=row_idx, column=col_idx, value=tot_g)
            c_g.number_format = '"₹ "#,##0.00" M"'
            c_g.font = Font(bold=True)
            c_g.fill = PatternFill(start_color="f1f5f9", fill_type="solid") # Slightly different for gross
            c_g.border = table_border
            c_g.alignment = Alignment(horizontal="right")
            
            row_idx += 1
    
        # Add Footer Rows in Excel
        ws_months.cell(row=row_idx, column=1, value="Grand Total (Net)").font = header_font
        ws_months.merge_cells(start_row=row_idx, start_column=1, end_row=row_idx, end_column=4)
        for c in range(1, 5):
            ws_months.cell(row=row_idx, column=c).fill = header_fill
            ws_months.cell(row=row_idx, column=c).border = table_border
        
        col_idx = 5
        m_totals_net = {}
        m_totals_gross = {}
        m_totals_returned = {}
        g_total_net = 0
        g_total_gross = 0
        g_total_returned = 0
        
        # Pre-calculate totals for footer
        for r in merged_data.values():
            g_total_net += r["total"]
            g_total_gross += r["total_gross"]
            for mk, mv in r["months"].items():
                m_totals_net[mk] = m_totals_net.get(mk, 0) + mv
                
        # We need to calculate monthly gross totals properly too
        for row in data:
            date_str = str(row.get("delivery_date") or row.get("invoice_date") or row.get("posting_date") or "")
            try:
                d = frappe.utils.getdate(date_str)
                m_key = d.strftime("%b %Y")
            except: m_key = "Unknown"
            
            gross_amt = flt(row.get("gross_amount") or 0)
            m_totals_gross[m_key] = m_totals_gross.get(m_key, 0) + gross_amt
            
            if row.get("is_return"):
                ret_val = abs(flt(row.get("amt_allocated") or 0))
                m_totals_returned[m_key] = m_totals_returned.get(m_key, 0) + ret_val
                g_total_returned += ret_val
    
        for m_key in sorted_months:
            v_net = flt(m_totals_net.get(m_key, 0))
            c = ws_months.cell(row=row_idx, column=col_idx, value=v_net)
            c.number_format = '"₹ "#,##0.00" M"'
            c.font = header_font
            c.fill = header_fill
            c.border = table_border
            c.alignment = Alignment(horizontal="right")
            col_idx += 1
            
        c_gn = ws_months.cell(row=row_idx, column=col_idx, value=g_total_net)
        c_gn.number_format = '"₹ "#,##0.00" M"'
        c_gn.font = header_font
        c_gn.fill = header_fill
        c_gn.border = table_border
        c_gn.alignment = Alignment(horizontal="right")
        col_idx += 1
        
        c_sep = ws_months.cell(row=row_idx, column=col_idx, value="-")
        c_sep.font = header_font
        c_sep.fill = header_fill
        c_sep.border = table_border
        c_sep.alignment = Alignment(horizontal="center")
        
        row_idx += 1
        
        # 3. Grand Total (Gross)
        ws_months.cell(row=row_idx, column=1, value="Grand Total (Gross)").font = header_font
        ws_months.merge_cells(start_row=row_idx, start_column=1, end_row=row_idx, end_column=4)
        for c in range(1, 5):
            ws_months.cell(row=row_idx, column=c).fill = header_fill
            ws_months.cell(row=row_idx, column=c).border = table_border
            
        col_idx = 5
        for m_key in sorted_months:
            v_gross = flt(m_totals_gross.get(m_key, 0))
            c = ws_months.cell(row=row_idx, column=col_idx, value=v_gross)
            c.number_format = '"₹ "#,##0.00" M"'
            c.font = header_font
            c.fill = header_fill
            c.border = table_border
            c.alignment = Alignment(horizontal="right")
            col_idx += 1
    
        c_sep2 = ws_months.cell(row=row_idx, column=col_idx, value="-")
        c_sep2.font = header_font
        c_sep2.fill = header_fill
        c_sep2.border = table_border
        c_sep2.alignment = Alignment(horizontal="center")
        col_idx += 1
        
        c_gg = ws_months.cell(row=row_idx, column=col_idx, value=g_total_gross)
        c_gg.number_format = '"₹ "#,##0.00" M"'
        c_gg.font = header_font
        c_gg.fill = header_fill
        c_gg.border = table_border
        c_gg.alignment = Alignment(horizontal="right")
        
        row_idx += 2

    if export_type in ["all", "detail"]:
        # -------------------------------------------------------------------------
        # Sheet 3: Sales Invoices List
        # -------------------------------------------------------------------------
        row_idx = 1
        ws_list.cell(row=row_idx, column=1, value="Detailed Sales Invoices List (Million INR)").font = section_font
        row_idx += 1
        
        ui_columns = [
            {"label": "S.No.", "fieldname": "sr_no_idx", "width": 8},
            {"label": "Invoice ID", "fieldname": "invoice_id", "width": 18},
            {"label": "Date", "fieldname": "invoice_date", "width": 14},
            {"label": "Type", "fieldname": "dom_exp", "width": 14},
            {"label": "Invoice Type", "fieldname": "invoice_type", "width": 20},
            {"label": "Status", "fieldname": "status", "width": 14},
            {"label": "Customer", "fieldname": "customer_name", "width": 25},
            {"label": "Business Region", "fieldname": "business_region_name", "width": 20},
            {"label": "Item", "fieldname": "item_code", "width": 20},
            {"label": "Sales Person", "fieldname": "sales_person", "width": 20},
            {"label": "Qty", "fieldname": "qty", "width": 10},
            {"label": "Amount (M)", "fieldname": "base_amount", "width": 18},
        ]
        
        for idx, col in enumerate(ui_columns, start=1):
            cell = ws_list.cell(row=row_idx, column=idx, value=col["label"])
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center")
            cell.border = table_border
            ws_list.column_dimensions[get_column_letter(idx)].width = col["width"]
        row_idx += 1
        
        for r_idx, row in enumerate(data):
            row_fill = zebra_fill if r_idx % 2 != 0 else None
            for idx, col in enumerate(ui_columns, start=1):
                fname = col["fieldname"]
                val = row.get(fname)
                
                if fname == "sr_no_idx":
                    val = r_idx + 1
                elif fname == "invoice_date":
                    val = row.get("delivery_date") or row.get("invoice_date") or row.get("posting_date")
                
                cell = ws_list.cell(row=row_idx, column=idx)
                cell.border = table_border
                if row_fill: cell.fill = row_fill
                
                if fname in ["qty", "base_amount"]:
                    num_val = flt(val or 0)
                    if fname == "base_amount":
                        num_val = flt(row.get("amt_allocated") or 0)
                        cell.number_format = '"₹ "#,##0.00" M"'
                    else:
                        cell.number_format = "#,##0.00"
                    cell.value = num_val
                    cell.alignment = Alignment(horizontal="right")
                else:
                    cell.value = str(val) if val is not None else ""
                    cell.alignment = Alignment(horizontal="left")
            row_idx += 1
        
        # -------------------------------------------------------------------------
        # Add Total Row for Detailed Invoice List
        # -------------------------------------------------------------------------
        ws_list.cell(row=row_idx, column=1, value="Grand Total").font = header_font
        ws_list.merge_cells(start_row=row_idx, start_column=1, end_row=row_idx, end_column=11)
        for c in range(1, 12):
            ws_list.cell(row=row_idx, column=c).fill = header_fill
            ws_list.cell(row=row_idx, column=c).border = table_border
    
        # Calculate Total Amount for the list
        total_list_amt = 0
        for row in data:
            total_list_amt += (row.get("amt_allocated") or 0)
    
        cell_total = ws_list.cell(row=row_idx, column=12, value=total_list_amt)
        cell_total.font = header_font
        cell_total.fill = header_fill
        cell_total.number_format = '"₹ "#,##0.00" M"'
        cell_total.alignment = Alignment(horizontal="right")
        cell_total.border = table_border

    # Remove dummy sheets if created
    for dummy_name in ["Dummy1", "Dummy2"]:
        if dummy_name in wb.sheetnames:
            wb.remove(wb[dummy_name])

    # Final Column Widths Adjustments
    if "Dashboard Overview" in wb.sheetnames:
        for i in range(1, 10):
            wb["Dashboard Overview"].column_dimensions[get_column_letter(i)].width = 20
    if "Month-Wise Revenue" in wb.sheetnames:
        for i in range(1, 10):
            wb["Month-Wise Revenue"].column_dimensions[get_column_letter(i)].width = 20

    # Save
    output = BytesIO()
    wb.save(output)
    output.seek(0)
    
    filename = f"Sales_Revenue_Dashboard_{frappe.utils.nowdate()}.xlsx"
    if export_type == "summary":
        filename = f"Month_Wise_Consolidated_Revenue_{frappe.utils.nowdate()}.xlsx"
    elif export_type == "detail":
        filename = f"Sales_Invoices_List_{frappe.utils.nowdate()}.xlsx"

    return {
        "filename": filename,
        "filecontent": base64.b64encode(output.read()).decode()
    }

