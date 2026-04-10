import frappe
from frappe import _
from frappe.utils import flt
from renu_customization.renu_customization.report.sales_invoice_report.sales_invoice_report import execute
import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from io import BytesIO
import base64
import json

def prepare_filters(filters):
    if isinstance(filters, str):
        filters = frappe.parse_json(filters)
    
    if not filters:
        filters = frappe._dict()

    # Handle DateRange from JS
    if filters.get("date_range"):
        date_range = filters.get("date_range")
        if isinstance(date_range, list) and len(date_range) == 2:
            filters["from_date"] = date_range[0]
            filters["to_date"] = date_range[1]
    
    return filters

@frappe.whitelist()
def get_dashboard_data(filters=None):
    filters = prepare_filters(filters)

    # Pass filters to report execute
    report_result = execute(filters)
    columns = report_result[0]
    data = report_result[1]
    
    if not data:
        return {
            "summary": [], "charts": {}, "results": [], "columns": columns
        }

    # Calculate Summaries (KPIs)
    total_rev = 0
    dom_rev = 0
    exp_rev = 0
    total_qty = 0
    unique_invoices = set()

    for row in data:
        amt = flt(row.get("base_amount"))
        total_rev += amt
        total_qty += flt(row.get("qty"))
        unique_invoices.add(row.get("invoice_id"))

        if row.get("dom_exp") == "Domestic":
            dom_rev += amt
        else:
            exp_rev += amt

    report_summary = [
        {"label": _("Total Revenue"), "value": total_rev, "indicator": "blue", "fieldtype": "Currency", "currency": "INR"},
        {"label": _("Domestic Revenue"), "value": dom_rev, "indicator": "green", "fieldtype": "Currency", "currency": "INR"},
        {"label": _("Export Revenue"), "value": exp_rev, "indicator": "orange", "fieldtype": "Currency", "currency": "INR"}
    ]

    # Aggregate specifically for dashboard charts requested
    # (Existing balance of aggregation logic remains same)
    sp_revenue = {}
    cust_revenue = {}
    prod_revenue = {}
    unique_items = {}
    
    prod_revenue = {}
    prod_names = {}
    unique_items = {}
    
    for row in data:
        inv_id = row.get("invoice_id")
        sp = row.get("sales_person") or "Unassigned"
        amt = flt(row.get("base_amount"))
        sp_revenue[sp] = sp_revenue.get(sp, 0) + amt
        
        item_key = f"{inv_id}_{row.get('item_code')}_{row.get('qty')}_{row.get('base_amount')}"
        if item_key not in unique_items:
            unique_items[item_key] = row
            cust = row.get("customer_name") or "Unknown"
            cust_revenue[cust] = cust_revenue.get(cust, 0) + amt
            
            prod_code = row.get("item_code") or "Unknown"
            prod_revenue[prod_code] = prod_revenue.get(prod_code, 0) + amt
            prod_names[prod_code] = row.get("item_name") or ""

    def get_chart_def(title, data_dict, label_key, limit=10):
        sorted_items = sorted(data_dict.items(), key=lambda x: x[1], reverse=True)
        top_items = sorted_items[:limit]
        return {
            "title": title,
            "data": {
                "labels": [x[0] for x in top_items],
                "datasets": [{"name": title, "values": [x[1] for x in top_items]}]
            },
            "type": "donut",
            "height": 300,
            "colors": ['#3498db', '#e74c3c', '#2ecc71', '#f1c40f', '#9b59b6', '#34495e', '#ecf0f1', '#1abc9c', '#d35400', '#7f8c8d']
        }

    def get_table_data(data_dict):
        sorted_items = sorted(data_dict.items(), key=lambda x: x[1], reverse=True)
        return [{"name": x[0], "value": x[1]} for x in sorted_items]

    return {
        "summary": report_summary,
        "charts": {
            "top_5_salesperson": get_chart_def("Top 5 Salesperson by Revenue", sp_revenue, "sales_person", limit=5),
            "top_10_customers": get_chart_def("Top 10 Customers by Revenue", cust_revenue, "customer", limit=10),
            "top_10_products": get_chart_def("Top 10 Products by Revenue", prod_revenue, "item_code", limit=10)
        },
        "tables": {
            "salesperson": get_table_data(sp_revenue),
            "customer": get_table_data(cust_revenue),
            "product": [{"code": k, "name": prod_names.get(k, ""), "value": v} for k, v in sorted(prod_revenue.items(), key=lambda x: x[1], reverse=True)]
        },
        "results": data,
        "columns": columns
    }

@frappe.whitelist()
def export_to_excel(filters=None, invoice_id=None):
    filters = prepare_filters(filters)
    if invoice_id:
        filters["invoice_id"] = invoice_id

    # Get data using report's execute function
    columns, data = execute(filters)

    if not data:
        return None

    # Create Workbook
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sales Invoice Data"

    # Header Info
    ws["A1"].value = "Revenue Dashboard Export"
    ws["A1"].font = Font(bold=True)
    ws["B1"].value = "Generated On: " + frappe.utils.now_datetime().strftime("%Y-%m-%d %H:%M:%S")

    # Column Headers
    row_idx = 3
    for idx, col in enumerate(columns, start=1):
        cell = ws.cell(row=row_idx, column=idx, value=col.get("label"))
        cell.font = Font(bold=True)
        cell.fill = PatternFill(start_color="D3D3D3", fill_type="solid")
        cell.alignment = Alignment(horizontal="center")

    # Data Rows
    row_idx += 1
    numeric_fields = ["qty", "item_rate", "amount", "base_amount", "exchange_rate", "item_purchase_rate"]
    
    for row in data:
        for idx, col in enumerate(columns, start=1):
            fieldname = col.get("fieldname")
            value = row.get(fieldname)
            cell = ws.cell(row=row_idx, column=idx)

            if fieldname in numeric_fields:
                cell.value = flt(value or 0)
                cell.number_format = "#,##0.00"
                cell.alignment = Alignment(horizontal="right")
            else:
                cell.value = str(value) if value is not None else ""
                cell.alignment = Alignment(horizontal="left")
        row_idx += 1

    # Adjust Column Widths
    for idx in range(1, len(columns) + 1):
        ws.column_dimensions[get_column_letter(idx)].width = 20

    # Save to buffer
    output = BytesIO()
    wb.save(output)
    output.seek(0)
    
    # Return base64 for JS to download
    return {
        "filename": f"Revenue_Dashboard_Sales_{frappe.utils.nowdate()}.xlsx",
        "filecontent": base64.b64encode(output.read()).decode()
    }
