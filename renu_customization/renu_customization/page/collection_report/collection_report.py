import frappe
from frappe import _
from frappe.utils import flt, nowdate, now_datetime
import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
from io import BytesIO
import base64

@frappe.whitelist()
def export_to_pdf(html):
    pdf_content = frappe.utils.pdf.get_pdf(html, {"orientation": "Landscape"})
    frappe.response.filename = f"Collection_Report_{nowdate()}.pdf"
    frappe.response.filecontent = pdf_content
    frappe.response.type = "download"

def prepare_filters(filters):
    if not filters:
        filters = {}
    elif isinstance(filters, str):
        filters = frappe.parse_json(filters)
    
    # Handle DateRange from JS
    if filters.get("date_range"):
        dr = filters.get("date_range")
        if isinstance(dr, list) and len(dr) == 2:
            filters["from_date"] = dr[0]
            filters["to_date"] = dr[1]

    # Handle Fiscal Year
    if filters.get("fiscal_year") and (not filters.get("from_date") or not filters.get("to_date")):
        fy = frappe.get_doc("Fiscal Year", filters.get("fiscal_year"))
        if fy:
            filters["from_date"] = filters.get("from_date") or fy.year_start_date
            filters["to_date"] = filters.get("to_date") or fy.year_end_date
    
    return frappe._dict(filters)

@frappe.whitelist()
def get_dashboard_data(filters=None):
    filters = prepare_filters(filters)

    conditions = [
        "pe.docstatus = 1",
        "pe.payment_type = 'Receive'",
        "pe.party_type = 'Customer'",
        "per.reference_doctype = 'Sales Invoice'"
    ]
    
    if filters.get("customer"):
        conditions.append(f"pe.party = {frappe.db.escape(filters.get('customer'))}")
    
    if filters.get("from_date"):
        conditions.append(f"pe.posting_date >= {frappe.db.escape(filters.get('from_date'))}")
    
    if filters.get("to_date"):
        conditions.append(f"pe.posting_date <= {frappe.db.escape(filters.get('to_date'))}")
 
    if filters.get("dom_exp"):
        if filters.get("dom_exp") == "Domestic":
            conditions.append("si.is_domestic = 1")
        elif filters.get("dom_exp") == "Export":
            conditions.append("si.is_export = 1")
 
    where_clause = " AND ".join(conditions)
    
    query = f"""
        SELECT 
            pe.name as payment_entry,
            pe.posting_date,
            pe.party as customer,
            per.reference_name as name,
            per.base_allocated_amount,
            si.is_export, 
            si.is_domestic,
            si.status,
            sii.item_code,
            sii.item_name,
            sii.base_amount as item_amount,
            si.base_net_total
        FROM `tabPayment Entry` pe
        JOIN `tabPayment Entry Reference` per ON per.parent = pe.name
        JOIN `tabSales Invoice` si ON si.name = per.reference_name
        JOIN `tabSales Invoice Item` sii ON sii.parent = si.name
        WHERE {where_clause}
        ORDER BY pe.posting_date DESC, pe.name DESC
    """
    
    data = frappe.db.sql(query, as_dict=True)
    
    # Fetch Sales Person from Sales Team child table
    inv_names = list(set([d.name for d in data]))
    sales_map = {}
    if inv_names:
        sales_team = frappe.get_all("Sales Team", 
            filters={"parent": ("in", inv_names), "parenttype": "Sales Invoice"},
            fields=["parent", "sales_person"]
        )
        for st in sales_team:
            if st.parent not in sales_map:
                sales_map[st.parent] = []
            sales_map[st.parent].append(st.sales_person)
    
    # Process data and apply sales person filter
    final_data = []
    sp_filter = filters.get("sales_person")
    
    # Track unique allocations for KPI calculations to avoid double counting due to item join
    unique_allocations = {}
    
    for d in data:
        d.sales_person = ", ".join(sales_map.get(d.name, []))
        
        if sp_filter and sp_filter not in (d.sales_person or ""):
            continue
        
        # Format Item
        d.item = f"{d.item_code} {d.item_name}" if d.item_code else (d.item_name or "")
        
        # Calculate item's share of the allocated payment (Pro-rata allocation)
        if flt(d.base_net_total) > 0:
            d.allocated_amount = (flt(d.item_amount) / flt(d.base_net_total)) * flt(d.base_allocated_amount)
        else:
            d.allocated_amount = flt(d.item_amount)
            
        final_data.append(d)
        
        # For KPIs, we sum unique payment-to-invoice allocations
        alloc_key = f"{d.payment_entry}-{d.name}"
        if alloc_key not in unique_allocations:
            unique_allocations[alloc_key] = {
                "amount": flt(d.base_allocated_amount),
                "is_export": d.is_export,
                "is_domestic": d.is_domestic
            }
 
    # Calculate KPIs from unique allocations - round each value to 4 decimal places in M
    total_collection = sum(round(v["amount"] / 1000000, 4) for v in unique_allocations.values()) * 1000000
    export_collection = sum(round(v["amount"] / 1000000, 4) for v in unique_allocations.values() if v["is_export"]) * 1000000
    domestic_collection = sum(round(v["amount"] / 1000000, 4) for v in unique_allocations.values() if not v["is_export"]) * 1000000
    
    summary = [
        {"label": _("TOTAL Collection"), "value": total_collection, "indicator": "Blue"},
        {"label": _("Export Collection"), "value": export_collection, "indicator": "Green"},
        {"label": _("Domestic Collection"), "value": domestic_collection, "indicator": "Orange"}
    ]
    
    # Pie Chart Data
    chart = {
        "title": _("Collection Breakdown (Export vs Domestic)"),
        "data": {
            "labels": [_("Export"), _("Domestic")],
            "datasets": [
                {
                    "name": _("Collection"),
                    "values": [export_collection, domestic_collection]
                }
            ]
        },
        "type": "donut",
        "colors": ["#10b981", "#f59e0b"]
    }
    
    return {
        "summary": summary,
        "chart": chart,
        "results": final_data
    }

@frappe.whitelist()
def export_to_excel(filters=None, export_type="all"):
    dashboard_data = get_dashboard_data(filters)
    data = dashboard_data.get("results")
    summary = dashboard_data.get("summary")
    
    if not data:
        return None

    wb = openpyxl.Workbook()
    
    # Based on export_type, we decide which sheets to keep
    if export_type == "all":
        ws_overview = wb.active
        ws_overview.title = "Dashboard Overview"
        ws_list = wb.create_sheet("Collection List")
    elif export_type == "detail":
        ws_list = wb.active
        ws_list.title = "Collection List"
        ws_overview = wb.create_sheet("Dummy") # Will be removed
    
    # Styling
    header_fill = PatternFill(start_color="2c3e50", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF")
    title_font = Font(bold=True, size=14)
    section_font = Font(bold=True, size=12)
    thin_side = Side(style='thin')
    table_border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)
    
    if export_type == "all":
        # 1. Overview Sheet
        ws_overview.cell(row=1, column=1, value="Collection Report Dashboard (Million INR)").font = title_font
        ws_overview.cell(row=1, column=4, value="Generated On: " + now_datetime().strftime("%Y-%m-%d %H:%M"))
        
        ws_overview.cell(row=3, column=1, value="Collection Metrics Summary").font = section_font
        
        colors = {"blue": "3b82f6", "green": "10b981", "orange": "f59e0b"}
        
        for i, s in enumerate(summary):
            r = 5 + (i // 3) * 3
            c = 1 + (i % 3) * 2
            bg_color = colors.get(s.get('indicator', 'blue').lower(), "3b82f6")
            
            cell_l = ws_overview.cell(row=r, column=c, value=s.get('label'))
            cell_l.font = Font(bold=True, color="FFFFFF")
            cell_l.fill = PatternFill(start_color=bg_color, fill_type="solid")
            cell_l.alignment = Alignment(horizontal="center")
            ws_overview.merge_cells(start_row=r, start_column=c, end_row=r, end_column=c+1)
            
            cell_v = ws_overview.cell(row=r+1, column=c, value=flt(s.get('value')) / 1000000)
            cell_v.font = Font(bold=True, size=11)
            cell_v.number_format = '"₹ "#,##0.0000" M"'
            cell_v.alignment = Alignment(horizontal="center")
            cell_v.border = Border(bottom=Side(style='medium', color=bg_color))
            ws_overview.merge_cells(start_row=r+1, start_column=c, end_row=r+1, end_column=c+1)

    if export_type in ["all", "detail"]:
        # 2. Collection List Sheet
        row_idx = 1
        ws_list.cell(row=row_idx, column=1, value="Detailed Collection List").font = section_font
        row_idx += 2
        
        headers = ["S.No.", "Invoice ID", "Date", "Customer", "Item", "Sales Person", "Amount (M)", "Type", "Status"]
        for idx, h in enumerate(headers, start=1):
            cell = ws_list.cell(row=row_idx, column=idx, value=h)
            cell.font, cell.fill, cell.alignment, cell.border = header_font, header_fill, Alignment(horizontal="center"), table_border
        row_idx += 1
        
        for r_idx, row in enumerate(data):
            ws_list.cell(row=row_idx, column=1, value=r_idx + 1).border = table_border
            ws_list.cell(row=row_idx, column=2, value=row['name']).border = table_border
            ws_list.cell(row=row_idx, column=3, value=row['posting_date']).border = table_border
            ws_list.cell(row=row_idx, column=4, value=row['customer']).border = table_border
            ws_list.cell(row=row_idx, column=5, value=row['item']).border = table_border
            ws_list.cell(row=row_idx, column=6, value=row['sales_person']).border = table_border
            
            amt_cell = ws_list.cell(row=row_idx, column=7, value=flt(row['allocated_amount']) / 1000000)
            amt_cell.number_format, amt_cell.border = '"₹ "#,##0.0000" M"', table_border
            
            ws_list.cell(row=row_idx, column=8, value="Export" if row['is_export'] else "Domestic").border = table_border
            ws_list.cell(row=row_idx, column=9, value=row['status']).border = table_border
            row_idx += 1
    
        # Add Total Row
        ws_list.cell(row=row_idx, column=1, value="Total").font = header_font
        ws_list.merge_cells(start_row=row_idx, start_column=1, end_row=row_idx, end_column=6)
        for c in range(1, 7):
            ws_list.cell(row=row_idx, column=c).fill = header_fill
            ws_list.cell(row=row_idx, column=c).border = table_border
            if c == 1:
                ws_list.cell(row=row_idx, column=c).alignment = Alignment(horizontal="right")
                
        total_amt = sum(flt(r['allocated_amount']) for r in data) / 1000000
        total_cell = ws_list.cell(row=row_idx, column=7, value=total_amt)
        total_cell.font = header_font
        total_cell.fill = header_fill
        total_cell.number_format, total_cell.border = '"₹ "#,##0.0000" M"', table_border
        
        ws_list.cell(row=row_idx, column=8, value="").fill = header_fill
        ws_list.cell(row=row_idx, column=8, value="").border = table_border
        ws_list.cell(row=row_idx, column=9, value="").fill = header_fill
        ws_list.cell(row=row_idx, column=9, value="").border = table_border
        row_idx += 1

    # Remove dummy if detail only
    if export_type == "detail" and "Dummy" in wb.sheetnames:
        wb.remove(wb["Dummy"])

    # Column Widths
    for ws in wb.worksheets:
        for i in range(1, 10):
            ws.column_dimensions[get_column_letter(i)].width = 20

    output = BytesIO()
    wb.save(output)
    output.seek(0)
    
    filename = f"Collection_Report_{nowdate()}.xlsx"
    if export_type == "detail":
        filename = f"Detailed_Collection_List_{nowdate()}.xlsx"
    
    return {
        "filename": filename,
        "filecontent": base64.b64encode(output.read()).decode()
    }
