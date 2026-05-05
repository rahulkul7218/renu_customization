import frappe
from frappe import _
from frappe.utils import flt, getdate, nowdate, date_diff, now_datetime
import json
import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
from io import BytesIO
import base64
from renu_customization.renu_customization.report.sales_invoice_report.sales_invoice_report import execute

@frappe.whitelist()
def export_to_pdf(html):
    pdf_content = frappe.utils.pdf.get_pdf(html, {"orientation": "Landscape"})
    frappe.response.filename = f"Margin_Build_{nowdate()}.pdf"
    frappe.response.filecontent = pdf_content
    frappe.response.type = "download"

@frappe.whitelist()
def get_dashboard_data(filters=None):
    if isinstance(filters, str):
        filters = json.loads(filters)
    if not filters:
        filters = {}

    # Handle Fiscal Year
    if filters.get("fiscal_year"):
        fy = frappe.get_doc("Fiscal Year", filters.get("fiscal_year"))
        if fy:
            filters["from_date"] = fy.year_start_date
            filters["to_date"] = fy.year_end_date

    base_filters = frappe._dict({})
    if filters.get("from_date"):
        base_filters.from_date = filters.get("from_date")
    if filters.get("to_date"):
        base_filters.to_date = filters.get("to_date")

    report_result = execute(base_filters)
    raw_data = report_result[1] if report_result and len(report_result) > 1 else []

    data = []
    
    inv_names = list(set([d.get("invoice_id") or d.get("name") or d.get("parent") for d in raw_data if d.get("invoice_id") or d.get("name") or d.get("parent")])) if raw_data else []
    dom_exp_map = {}
    status_map = {}
    
    if inv_names:
        invoices = frappe.get_all("Sales Invoice", filters={"name": ("in", inv_names)}, fields=["name", "is_domestic", "is_export", "invoice_type", "status"])
        for i in invoices:
            status_map[i.name] = i.status
            if i.is_domestic: 
                dom_exp_map[i.name] = "Domestic"
            elif i.is_export: 
                dom_exp_map[i.name] = "Export"
            elif i.invoice_type and "Domestic" in i.invoice_type:
                dom_exp_map[i.name] = "Domestic"
            elif i.invoice_type and "Export" in i.invoice_type:
                dom_exp_map[i.name] = "Export"
            else: 
                dom_exp_map[i.name] = "Other"

    for row in raw_data:
        inv_id = row.get("invoice_id") or row.get("name") or row.get("parent")
        
        status = status_map.get(inv_id)
        if status == "Cancelled":
            continue

        revenue = flt(row.get("base_amount"))
        qty = flt(row.get("qty"))
        purchase_rate = flt(row.get("item_purchase_rate"))
        cogs = qty * purchase_rate
        margin = revenue - cogs
        
        row_type = dom_exp_map.get(inv_id, "Other")

        if filters.get("type") and row_type != filters.get("type"):
            continue

        if filters.get("customer") and filters.get("customer") != row.get("customer"):
            continue
            
        data.append({
            "name": inv_id,
            "posting_date": row.get("invoice_date") or row.get("posting_date"),
            "customer": row.get("customer"),
            "customer_name": row.get("customer_name"),
            "item_code": row.get("item_code"),
            "item_name": row.get("item_name"),
            "item": f"{row.get('item_code')} {row.get('item_name')}" if row.get('item_code') else (row.get('item_name') or ""),
            "type": row_type,
            "margin": margin,
            "revenue": revenue,
            "cogs": cogs
        })

    # No longer aggregating by invoice to show items
    final_results = data
    final_results.sort(key=lambda x: getdate(x["posting_date"]) if x["posting_date"] else getdate(), reverse=True)

    total_margin = 0
    export_margin = 0
    domestic_margin = 0
    
    unique_invoices = set()

    for d in final_results:
        amt = flt(d["margin"])
        total_margin += amt
        if d["type"] == "Export":
            export_margin += amt
        elif d["type"] == "Domestic":
            domestic_margin += amt
        
        unique_invoices.add(d["name"])

    summary = [
        {"label": _("Total Margin"), "value": total_margin, "indicator": "green", "fieldtype": "Currency"},
        {"label": _("Export Margin"), "value": export_margin, "indicator": "orange", "fieldtype": "Currency"},
        {"label": _("Domestic Margin"), "value": domestic_margin, "indicator": "blue", "fieldtype": "Currency"},
        {"label": _("Total Invoices"), "value": len(unique_invoices), "indicator": "purple", "fieldtype": "Int"}
    ]

    charts = {
        "margin_breakdown": {
            "title": _("Margin Breakdown (Export vs Domestic)"),
            "data": {
                "labels": [_("Export Margin"), _("Domestic Margin")],
                "datasets": [{"name": _("Margin"), "values": [export_margin, domestic_margin]}]
            },
            "type": "donut",
            "colors": ["#f59e0b", "#3b82f6"]
        }
    }

    return {
        "summary": summary,
        "charts": charts,
        "results": final_results
    }

@frappe.whitelist()
def export_to_excel(filters=None, export_type="all"):
    dashboard_data = get_dashboard_data(filters)
    data = dashboard_data.get("results")
    summary = dashboard_data.get("summary")
    
    if not data:
        return None

    wb = openpyxl.Workbook()
    
    # Based on export_type, we decide which sheets to keep
    if export_type == "all":
        ws_overview = wb.active
        ws_overview.title = "Dashboard Overview"
        ws_list = wb.create_sheet("Margin List")
    elif export_type == "detail":
        ws_list = wb.active
        ws_list.title = "Margin List"
        ws_overview = wb.create_sheet("Dummy") # Will be removed
    
    header_fill = PatternFill(start_color="2c3e50", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF")
    title_font = Font(bold=True, size=14)
    section_font = Font(bold=True, size=12)
    thin_side = Side(style='thin')
    table_border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)
    
    if export_type == "all":
        ws_overview.cell(row=1, column=1, value="Margin Build Dashboard (Million INR)").font = title_font
        ws_overview.cell(row=1, column=4, value="Generated On: " + now_datetime().strftime("%Y-%m-%d %H:%M"))
        
        ws_overview.cell(row=3, column=1, value="Margin Metrics Summary").font = section_font
        
        colors = {"red": "ef4444", "orange": "f59e0b", "blue": "3b82f6", "green": "10b981", "purple": "8b5cf6"}
        
        for i, s in enumerate(summary):
            r = 5 + (i // 2) * 3
            c = 1 + (i % 2) * 2
            bg_color = colors.get(s.get('indicator', 'blue').lower(), "3b82f6")
            
            cell_l = ws_overview.cell(row=r, column=c, value=s.get('label'))
            cell_l.font = Font(bold=True, color="FFFFFF")
            cell_l.fill = PatternFill(start_color=bg_color, fill_type="solid")
            cell_l.alignment = Alignment(horizontal="center")
            ws_overview.merge_cells(start_row=r, start_column=c, end_row=r, end_column=c+1)
            
            val = s.get('value')
            if s.get('fieldtype') == 'Currency':
                cell_v = ws_overview.cell(row=r+1, column=c, value=flt(val) / 1000000)
                cell_v.number_format = '"₹ "#,##0.0000" M"'
            else:
                cell_v = ws_overview.cell(row=r+1, column=c, value=val)
            cell_v.font = Font(bold=True, size=11)
            cell_v.alignment = Alignment(horizontal="center")
            cell_v.border = Border(bottom=Side(style='medium', color=bg_color))
            ws_overview.merge_cells(start_row=r+1, start_column=c, end_row=r+1, end_column=c+1)

    if export_type in ["all", "detail"]:
        row_idx = 1
        ws_list.cell(row=row_idx, column=1, value="Detailed Margin List").font = section_font
        row_idx += 2
        
        headers = ["S.No.", "Invoice ID", "Date", "Customer", "Item", "Type", "Revenue (M)", "COGS (M)", "Margin (M)"]
        for idx, h in enumerate(headers, start=1):
            cell = ws_list.cell(row=row_idx, column=idx, value=h)
            cell.font, cell.fill, cell.alignment, cell.border = header_font, header_fill, Alignment(horizontal="center"), table_border
        row_idx += 1
        
        for r_idx, row in enumerate(data):
            ws_list.cell(row=row_idx, column=1, value=r_idx + 1).border = table_border
            ws_list.cell(row=row_idx, column=2, value=row['name']).border = table_border
            ws_list.cell(row=row_idx, column=3, value=row['posting_date']).border = table_border
            ws_list.cell(row=row_idx, column=4, value=row['customer_name'] or row['customer']).border = table_border
            ws_list.cell(row=row_idx, column=5, value=row['item']).border = table_border
            ws_list.cell(row=row_idx, column=6, value=row['type']).border = table_border
            
            rev_cell = ws_list.cell(row=row_idx, column=7, value=flt(row['revenue']) / 1000000)
            rev_cell.number_format, rev_cell.border = '"₹ "#,##0.0000" M"', table_border
            
            cogs_cell = ws_list.cell(row=row_idx, column=8, value=flt(row['cogs']) / 1000000)
            cogs_cell.number_format, cogs_cell.border = '"₹ "#,##0.0000" M"', table_border
            
            margin_cell = ws_list.cell(row=row_idx, column=9, value=flt(row['margin']) / 1000000)
            margin_cell.number_format, margin_cell.border = '"₹ "#,##0.0000" M"', table_border
            
            row_idx += 1
    
        # Add Grand Total Row for Margin List
        ws_list.cell(row=row_idx, column=1, value="Grand Total").font = header_font
        ws_list.merge_cells(start_row=row_idx, start_column=1, end_row=row_idx, end_column=6)
        for c in range(1, 7):
            ws_list.cell(row=row_idx, column=c).fill = header_fill
            ws_list.cell(row=row_idx, column=c).border = table_border
            
        total_rev = sum(flt(row['revenue']) for row in data) / 1000000
        c_rev = ws_list.cell(row=row_idx, column=7, value=total_rev)
        c_rev.number_format, c_rev.font, c_rev.fill, c_rev.border = '"₹ "#,##0.0000" M"', header_font, header_fill, table_border
        c_rev.alignment = Alignment(horizontal="right")
        
        total_cogs = sum(flt(row['cogs']) for row in data) / 1000000
        c_cogs = ws_list.cell(row=row_idx, column=8, value=total_cogs)
        c_cogs.number_format, c_cogs.font, c_cogs.fill, c_cogs.border = '"₹ "#,##0.0000" M"', header_font, header_fill, table_border
        c_cogs.alignment = Alignment(horizontal="right")
        
        total_margin = sum(flt(row['margin']) for row in data) / 1000000
        c_margin = ws_list.cell(row=row_idx, column=9, value=total_margin)
        c_margin.number_format, c_margin.font, c_margin.fill, c_margin.border = '"₹ "#,##0.0000" M"', header_font, header_fill, table_border
        c_margin.alignment = Alignment(horizontal="right")

    # Remove dummy if detail only
    if export_type == "detail" and "Dummy" in wb.sheetnames:
        wb.remove(wb["Dummy"])

    # Column Widths
    for ws in wb.worksheets:
        for i in range(1, 10):
            ws.column_dimensions[get_column_letter(i)].width = 25

    output = BytesIO()
    wb.save(output)
    output.seek(0)
    
    filename = f"Margin_Build_{nowdate()}.xlsx"
    if export_type == "detail":
        filename = f"Detailed_Margin_List_{nowdate()}.xlsx"
        
    return {
        "filename": filename,
        "filecontent": base64.b64encode(output.read()).decode()
    }
